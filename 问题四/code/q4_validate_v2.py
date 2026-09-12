"""Independent ledger/XLSX/provenance acceptance; no robustness reruns."""

from pathlib import Path

import json,hashlib,subprocess,sys,xml.etree.ElementTree as ET

import numpy as np

import pandas as pd

import openpyxl

from q4_data import load_q4_inputs,ETA_C,ETA_D,SOC_MIN,SOC_MAX,XMAX

from q4_revision_run import CODE,ROOT,REV,RUNS,VERSION,sha,load,dump,fingerprint

from q4_select import promotion_gate



class Suite:

    def __init__(self):self.rows=[]

    def check(self,key,name,ok,evidence):

        self.rows.append({'id':key,'name':name,'status':'PASS' if bool(ok) else 'FAIL','evidence':evidence})

        print(key,'PASS' if ok else 'FAIL',name,flush=True)

    def skip(self,key,name,why):self.rows.append({'id':key,'name':name,'status':'SKIP','evidence':{'reason':why}})



def maxabs(v):return float(np.max(np.abs(np.asarray(v,float))))



def xlsx_check(br,slots,ctr,days):

    fn='result4-2.xlsx' if br=='q4_2' else 'result4-3.xlsx';path=CODE.parent/fn

    wb=openpyxl.load_workbook(path,read_only=True,data_only=True);errs={};r2=ctr.sort_values(['day_index','plan_j'])

    for name,field,costfield in [('计划购电量','B_kwh','plan_reference_fee_yuan')]+([('调整购电量','A_kwh','final_regular_fee_yuan')] if br=='q4_3' else []):

        ws=wb[name];a=np.asarray(list(ws.iter_rows(min_row=2,max_row=335,min_col=2,max_col=147,values_only=True)),float)

        wanted=r2[field].to_numpy().reshape(334,144)

        errs[name+'_all_slots']=maxabs(a[:,:144]-wanted)

        errs[name+'_totals']=maxabs(a[:,144]-wanted.sum(1))

        errs[name+'_fees']=maxabs(a[:,145]-r2[costfield].to_numpy().reshape(334,144).sum(1))

        if (ws.max_row,ws.max_column)!=(335,147):errs[name+'_bad_geometry']=1

    a=list(wb['充放电量'].iter_rows(min_row=2,values_only=True));errs['storage_rows']=abs(len(a)-334*6)

    wanted=slots[['x','y']].to_numpy().reshape(334,6,24,2).sum(2).reshape(334*6,2)

    errs['storage_energy']=maxabs(np.asarray([[r[2],r[3]] for r in a],float)-wanted)

    errs['storage_S00']=maxabs(np.array([a[6*i][5] for i in range(334)])-days.S00_kwh)

    errs['storage_S24']=maxabs(np.array([a[6*i+1][5] for i in range(334)])-days.S24_kwh)

    actual={};ds=None

    for row in wb['紧急购电量'].iter_rows(min_row=2,values_only=True):

        if row[0] is not None:ds=row[0].date().isoformat()

        actual[ds]=actual.get(ds,0.)+float(row[2] or 0.)

    errs['emergency_by_day']=maxabs(np.array([actual[k] for k in days.date])-days.emergency_kwh)

    wb.close()

    return {'all_numeric_errors':errs,'sha256':sha(path),'path':str(path)},max(errs.values())<1e-5





def run(require_semantics=True):

    s=Suite();data=load_q4_inputs(ROOT);fr=load(CODE/'january_frozen_selection.json');prot=load(REV/'protected_hashes.json')

    changed=[n for n,h in prot['files'].items() if not (ROOT/n).is_file() or sha(ROOT/n)!=h]

    s.check('T42','Q1/Q2/Q3和官方附件内容保护',not changed,{'protected_files':len(prot['files']),'changed':changed,'base_q3_head':prot['base_q3_head'],'head_may_advance_after_review':True})

    cp=subprocess.run(['git','-c','core.quotepath=false','status','--porcelain=v1','-z'],cwd=ROOT,capture_output=True)

    entries=[z.decode('utf-8') for z in cp.stdout.split(b'\0') if z];outside=[z for z in entries if not z[3:].replace('\\','/').startswith('问题四/')]

    tracked_outside=[z for z in outside if not z.startswith('?? ')];untracked_outside=[z for z in outside if z.startswith('?? ')]

    ext=load(REV/'outside_scope_observation.json');preserved=(ROOT/ext['path']).is_file() and sha(ROOT/ext['path'])==ext['sha256']

    unexplained=[z for z in untracked_outside if z[3:]!=ext['path']]

    s.check('V01','已跟踪上游不变，范围外新增文件单列保留',cp.returncode==0 and not tracked_outside and not unexplained and preserved,{'tracked_changes_outside_q4':tracked_outside,'untracked_outside_q4':untracked_outside,'preserved_observation':ext,'preserved_hash_matches':preserved,'status_entries':len(entries)})

    s.check('V02','官方时段、维度与缺失首段',data.price_plan.shape==(365,144) and np.isnan(data.price_cal[0,0]) and data.price_cal[31,0]==data.price_plan[30,143],{'price_shape':list(data.price_plan.shape),'headers':[data.plan_headers[0],data.plan_headers[-1]],'initial_missing':bool(np.isnan(data.price_cal[0,0]))})

    junit=ET.parse(REV/'unit_tests.xml').getroot();ts=list(junit.iter('testsuite'));n=sum(int(t.attrib.get('tests',0)) for t in ts);failed=sum(int(t.attrib.get('failures',0))+int(t.attrib.get('errors',0)) for t in ts)

    s.check('V03','因果性、可行域嵌套与非凸结算单测',n>=19 and failed==0,{'tests':n,'failures_or_errors':failed,'junit_sha256':sha(REV/'unit_tests.xml')})

    artifact_hashes={};rows={}

    for br in ('q4_2','q4_3'):

        prefix=br+'-';p=CODE/br;sl=pd.read_csv(p/'physical_10min.csv');ctr=pd.read_csv(p/'contract_ledger.csv');dd=pd.read_csv(p/'daily_ledger.csv');ev=pd.read_csv(p/'event_ledger.csv');m=load(p/'formal_summary.json');tail=load(p/'tail_bridge.json');rows[br]=(sl,ctr,dd,ev,m)

        bi=load(REV/f'bridge_{br}.json');s.check(prefix+'I','365日逐段回退到最新Q2/Q3',bi['status']=='PASS' and bi['numerical_core_hash']==fingerprint() and max(bi['max_absolute_errors'].values())<1e-6,bi)

        s.check(prefix+'C','正式期覆盖与事件权限',len(sl)==48096 and len(ctr)==48096 and len(dd)==334 and len(ev)==334*(1 if br=='q4_2' else 4) and list(dd.day_index)==list(range(31,365)),{'slots':len(sl),'contracts':len(ctr),'days':len(dd),'events':len(ev)})

        current={n:sha(p/n) for n in m['result_sha256']};prov=m['provenance'];freeze=fr[br]

        s.check(prefix+'P','冻结配置/源代码/结果哈希',current==m['result_sha256'] and prov['numerical_core_hash']==fingerprint() and m['january_decision_sha256']==sha(RUNS/f'{br}_delivery_january/decision.json') and prov['core_file_sha256']=={n:sha(CODE/n) for n in prov['core_file_sha256']}, {'current_core_hash':fingerprint(),'run_core_hash':prov['numerical_core_hash'],'result_sha256':current})

        gate=promotion_gate(freeze['B1']);expected='B1' if gate['promote'] else 'B0'

        s.check(prefix+'F','一月统一晋级，不按正式期反选',m['selected']==freeze['selected']==expected and freeze['formal_period_examined'] is False and freeze['numerical_core_hash']==fingerprint(),{'selected':expected,'gate':gate,'January_cash_B0':freeze['B0']['cash_total_yuan'],'January_cash_B1':freeze['B1']['cash_total_yuan']})

        expected_cfg=dict(freeze['config']);actual_cfg=dict(m['config']);expected_cfg.pop('label',None);actual_cfg.pop('label',None)
        legal=all(all(int(h)<=int(row.day_index)-3 for h in json.loads(row.source_days)) for row in ev.itertuples())
        weights_ok=all(abs(sum(json.loads(row.scenario_weights))-1)<1e-10 and min(json.loads(row.scenario_weights))>=0 and len(json.loads(row.scenario_weights))==row.scenario_k for row in ev.itertuples())
        s.check(prefix+'INFO','冻结配置、已释放场景源日与权重',expected_cfg==actual_cfg and legal and weights_ok and float(ev.max_eq_residual.max())<1e-5 and m['formal_state_start']==freeze['state_end'],{'config_same_except_label':expected_cfg==actual_cfg,'source_day_release':legal,'weights':weights_ok,'max_planning_residual':float(ev.max_eq_residual.max())})
        residuals={'contract':maxabs(sl.qL+sl.qB+sl.u-sl.Q_kwh),'PV':maxabs(sl.gL+sl.gB+sl.kappa-sl.pv_kwh),'load':maxabs(sl.qL+sl.gL+sl.y+sl.eL-sl.load_kwh),'charge':maxabs(sl.x-sl.qB-sl.gB),'SOC':maxabs(sl.S1_kwh-sl.S0_kwh-ETA_C*sl.x+sl.y/ETA_D),'grid':maxabs(sl.I_kwh-sl.qL-sl.qB)}

        nonneg=float(sl[['qL','qB','u','gL','gB','kappa','eL','eB','x','y']].min().min())

        s.check(prefix+'PHY','独立重算全部源流方程',max(residuals.values())<1e-6 and nonneg>=-1e-8,{'residuals':residuals,'min_flow':nonneg})

        s.check(prefix+'S','功率、SOC、无同时充放电、禁止紧急充电',sl.S1_kwh.min()>=SOC_MIN-1e-6 and sl.S1_kwh.max()<=SOC_MAX+1e-6 and max(sl.x.max(),sl.y.max())<=XMAX+1e-6 and not ((sl.x>1e-7)&(sl.y>1e-7)).any() and maxabs(sl.eB)==0,{'soc_min':float(sl.S1_kwh.min()),'soc_max':float(sl.S1_kwh.max()),'max_charge':float(sl.x.max()),'max_discharge':float(sl.y.max()),'emergency_charge_max':float(sl.eB.max())})

        j=sl.slot.to_numpy(int);d=sl.day_index.to_numpy(int)

        s.check(prefix+'A','真实附件按自然时段逐条匹配',maxabs(sl.price_yuan_per_kwh-data.price_cal[d,j])<1e-12 and maxabs(sl.load_kwh-data.load_cal_kwh[d,j])<1e-8 and maxabs(sl.pv_kwh-data.pv_cal_kwh[d,j])<1e-8,{'checked':len(sl)})

        fee=sl.price_yuan_per_kwh*np.minimum(sl.B_kwh,sl.A_kwh)+.5*sl.price_yuan_per_kwh*np.maximum(sl.B_kwh-sl.A_kwh,0)+1.5*sl.price_yuan_per_kwh*np.maximum(sl.A_kwh-sl.B_kwh,0)

        errors={'regular':maxabs(fee-sl.regular_fee_yuan),'emergency':maxabs(5*sl.price_yuan_per_kwh*sl.e_kwh-sl.emergency_fee_yuan),'cash':maxabs(fee+sl.emergency_fee_yuan-sl.cash_fee_yuan),'daily':maxabs(sl.groupby('date').cash_fee_yuan.sum().to_numpy()-dd.cash_fee_yuan),'summary':abs(float(sl.cash_fee_yuan.sum())-m['formal']['cash_total_yuan'])}

        s.check(prefix+'M','原始B锚定、5倍紧急电、现金对账',max(errors.values())<1e-5,errors)

        cont=maxabs(sl.S1_kwh.to_numpy()[:-1]-sl.S0_kwh.to_numpy()[1:]);fstart=m['formal_state_start'];natplan=ctr.final_regular_fee_yuan.sum()-tail['regular_fee_yuan']+float(sl.regular_fee_yuan.iloc[0])

        s.check(prefix+'B','跨日SOC、2月首段与年末+1尾桥',cont<1e-8 and abs(sl.S0_kwh.iloc[0]-fstart['soc'])<1e-8 and abs(sl.A_kwh.iloc[0]-fstart['prev_A'][143])<1e-8 and abs(natplan-sl.regular_fee_yuan.sum())<1e-5 and tail['physics_residual']<1e-6 and abs(tail['S0_kwh']-sl.S1_kwh.iloc[-1])<1e-8,{'continuity_error':cont,'natural_plan_regular_difference':float(natplan-sl.regular_fee_yuan.sum()),'tail_S0':tail['S0_kwh'],'tail_residual':tail['physics_residual']})

        xe,ok=xlsx_check(br,sl,ctr,dd);s.check(prefix+'X','官方Excel全单元格独立重读',ok,xe);artifact_hashes[xe['path']]=xe['sha256']

        matched=load(RUNS/f'{br}_delivery_matched_fixed_decision/formal_summary.json')

        s.check(prefix+'MATCH','匹配固定价决策基线重新优化',matched['row_counts']['days']==334 and matched['config']['decision_fixed_price'] and matched['provenance']['numerical_core_hash']==fingerprint() and matched['formal_state_start']==m['formal_state_start'],{'formal_cash':m['formal']['cash_total_yuan'],'matched_cash':matched['formal']['cash_total_yuan'],'same_initial_state':True})

    if require_semantics:
        sem=load(RUNS/'q4_3_adjustment_time_selected/formal_summary.json');semjan=load(RUNS/'q4_3_adjustment_time_january/decision.json');semctr=pd.read_csv(RUNS/'q4_3_adjustment_time_selected/contract_ledger.csv');same=maxabs(semctr.A_kwh-rows['q4_3'][1].A_kwh)
        s.check('V04','调整事件价格的一月与334日真重优化',sem['row_counts']['days']==334 and semjan['formal_period_examined'] is False and sem['config']['settlement_clock']=='adjustment_time' and sem['provenance']['numerical_core_hash']==fingerprint() and same>1e-6 and sem['provenance'].get('semantic_adapter_sha256')==sha(CODE/'q4_semantics_fast.py') and semjan['provenance'].get('semantic_adapter_sha256')==sha(CODE/'q4_semantics_fast.py'),{'selected':sem['selected'],'cash':sem['formal']['cash_total_yuan'],'max_contract_change_not_repricing':same,'january_hash':sha(RUNS/'q4_3_adjustment_time_january/decision.json')})
        semdir=RUNS/'q4_3_adjustment_time_selected';asl=pd.read_csv(semdir/'physical_10min.csv');add=pd.read_csv(semdir/'daily_ledger.csv');aev=pd.read_csv(semdir/'event_ledger.csv');atail=load(semdir/'tail_bridge.json')
        ap=asl.price_yuan_per_kwh;ape=asl.adjustment_price_yuan_per_kwh;B=asl.B_kwh;A=asl.A_kwh
        independent_regular=ap*np.minimum(B,A)+.5*ape*np.maximum(B-A,0)+1.5*ape*np.maximum(A-B,0)
        alt_errors={'regular':maxabs(independent_regular-asl.regular_fee_yuan),'emergency':maxabs(5*ap*asl.e_kwh-asl.emergency_fee_yuan),'cash':maxabs(independent_regular+5*ap*asl.e_kwh-asl.cash_fee_yuan),'daily_cash':maxabs(asl.groupby('date').cash_fee_yuan.sum().to_numpy()-add.cash_fee_yuan),'summary_cash':abs(float(asl.cash_fee_yuan.sum())-sem['formal']['cash_total_yuan']),'contract':maxabs(asl.qL+asl.qB+asl.u-asl.Q_kwh),'PV':maxabs(asl.gL+asl.gB+asl.kappa-asl.pv_kwh),'load':maxabs(asl.qL+asl.gL+asl.y+asl.eL-asl.load_kwh),'SOC':maxabs(asl.S1_kwh-asl.S0_kwh-ETA_C*asl.x+asl.y/ETA_D),'SOC_continuity':maxabs(asl.S1_kwh.to_numpy()[:-1]-asl.S0_kwh.to_numpy()[1:])}
        si=asl.slot.to_numpy(int);di=asl.day_index.to_numpy(int);expected_p=data.price_cal[di,si].copy()
        expected_p[si==0]=data.price_cal[di[si==0]-1,108]
        for start in (36,72,108):
            mask=(si>=start)&(si<start+36);expected_p[mask]=data.price_cal[di[mask],start]
        alt_errors['event_price_clock']=maxabs(expected_p-ape)
        alt_errors['tail_bridge']=abs(float(semctr.final_regular_fee_yuan.sum()-atail['regular_fee_yuan']+asl.regular_fee_yuan.iloc[0]-asl.regular_fee_yuan.sum()))
        unchanged={n:sha(semdir/n) for n in sem['result_sha256']}==sem['result_sha256']
        s.check('V04B','替代语义独立源流/事件电价/费用与求解间隙复核',max(alt_errors.values())<1e-5 and unchanged and float(aev.mip_gap.max())<=1.0001e-7 and float(aev.max_eq_residual.max())<1e-6 and float(asl.eB.max())==0,{'max_errors':alt_errors,'max_certified_mip_gap':float(aev.mip_gap.max()),'max_constraint_residual':float(aev.max_eq_residual.max()),'result_hashes_match':unchanged})
    else:
        s.skip('V04','替代结算仍在运行','仅为主结果预验收；本文件不是最终验收。')
    fm=load(CODE/'figure_manifest.json');missing=[n for n,h in fm.items() if sha(CODE.parent/'output/figures'/n)!=h]

    s.check('V05','当前图表哈希',not missing and len(fm)>0,{'files':len(fm),'mismatches':missing})

    metadata=load(CODE/'data_manifest.json');public_hashes={n:sha(CODE.parent/n) for n in ('result4-2.xlsx','result4-3.xlsx')}
    s.check('V06','公开元数据与论文只指向新正式结果',metadata.get('backend')==VERSION and metadata['template_output_hashes']==public_hashes and all(VERSION in (CODE.parent/p).read_text(encoding='utf-8') for p in ('问题4论文材料.md','output/06_结果分析与风险对比.md')),{'public_result_hashes':public_hashes,'metadata_backend':metadata.get('backend')})
    s.skip('S01','DRO/参数/历史窗/效率等鲁棒性','用户明确不跑；旧V1.2证据已标STALE，不并入本轮通过数。')

    s.skip('S02','旧E0—E15消融矩阵','仅本轮明确重算的同源对照可引用；未重算条目不继承PASS。')

    s.skip('S03','旧Price Oracle','旧初态/控制器不匹配，不能沿用数值或宣称严格VOI。')

    s.skip('S04','Full-information Oracle','本轮不重跑，不保留新版严格下界数值或最优性差额主张。')

    report={'suite':VERSION+' correctness-only','all_executed_pass':all(r['status']!='FAIL' for r in s.rows),'passed':sum(r['status']=='PASS' for r in s.rows),'failed':sum(r['status']=='FAIL' for r in s.rows),'skipped':sum(r['status']=='SKIP' for r in s.rows),'unit_tests':n,'numerical_core_hash':fingerprint(),'tests':s.rows}

    if not require_semantics:
        dump(REV/'prevalidation.json',report)
        print(json.dumps({k:v for k,v in report.items() if k!='tests'},ensure_ascii=False),flush=True)
        if report['failed']:raise AssertionError('Main prevalidation failed')
        return report
    dump(CODE/'validation.json',report)

    head=subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip()

    manifest={'backend':VERSION,'base_q3_head':prot['base_q3_head'],'run_head':head,'working_tree_is_uncommitted':bool(entries),'official_input_hashes':{k:v for k,v in prot['files'].items() if k.startswith('26C题/')},'protected_snapshot_sha256':sha(REV/'protected_hashes.json'),'numerical_core_hash':fingerprint(),'model_config_hash':sha(CODE/'config_frozen.yaml'),'frozen_selection_sha256':sha(CODE/'january_frozen_selection.json'),'result4_hashes':artifact_hashes,'validation_sha256':sha(CODE/'validation.json'),'source_hashes':{p.name:sha(p) for p in CODE.glob('*.py')},'executed_validation_pass':report['all_executed_pass'],'skipped':['robustness','unrerun ablations','legacy Oracle','FI Oracle']}

    manifest['human_material_sha256']={str(p.relative_to(CODE.parent)).replace('\\','/'):sha(p) for p in [CODE.parent/'问题4论文材料.md',CODE/'README.md',*(CODE.parent/'output').glob('*.md')]}
    dump(CODE/'run_manifest.json',manifest)
    dump(CODE/'final_acceptance.json',{'backend':VERSION,'status':'NUMERICAL_PASS_PENDING_FINAL_VISUAL' if report['all_executed_pass'] else 'FAIL','passed':report['passed'],'failed':report['failed'],'skipped':report['skipped'],'unit_tests':n,'validation_sha256':sha(CODE/'validation.json'),'run_manifest_sha256':sha(CODE/'run_manifest.json'),'robustness_run':False})

    print(json.dumps({k:v for k,v in report.items() if k!='tests'},ensure_ascii=False),flush=True)

    if report['failed']:raise AssertionError('Acceptance failed; inspect validation.json')

    return report

if __name__=='__main__':run()

