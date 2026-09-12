# -*- coding: utf-8 -*-
"""Q4 official-data loader and the two time-coordinate systems.

The official 144-column plan row is [00:10-00:20, ..., 23:50-00:00+1,
00:00-00:10+1].  Natural-day slot i is [00:00+i*10min, ...].  Thus natural
slot 0 belongs to the previous plan row's last column, while natural slots
1..143 belong to the current plan row columns 0..142.  This module never
fills 2025-01-01 00:00-00:10 with a fabricated zero.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Dict, List, Optional, Sequence

import numpy as np
import openpyxl

DT = 1.0 / 6.0
CAP = 12000.0
PMAX = 5000.0
XMAX = PMAX * DT
SOC0 = 6000.0
SOC_MIN = 1200.0
SOC_MAX = 10800.0
ETA_C = float(np.sqrt(0.9))
ETA_D = float(np.sqrt(0.9))
EVENT_HOURS = (0, 6, 12, 18)
EVENT_PLAN_START = {0: 0, 6: 35, 12: 71, 18: 107}
FORMAL_START_DAY = 31  # 2025-02-01
FORMAL_END_DAY = 365   # end-exclusive, through 2025-12-31


@dataclass(frozen=True)
class ForecastRecord:
    issue_time: datetime
    target_time: datetime
    lead_hour: int
    pv_kw: float


@dataclass
class Q4Data:
    dates: List[datetime]
    plan_headers: List[str]
    load_plan_kwh: np.ndarray
    pv_plan_kwh: np.ndarray
    net_plan_kwh: np.ndarray
    load_cal_kwh: np.ndarray
    pv_cal_kwh: np.ndarray
    net_cal_kwh: np.ndarray
    price_plan: np.ndarray
    price_cal: np.ndarray
    fixed_price_plan: np.ndarray
    fixed_price_cal: np.ndarray
    forecasts: Dict[datetime, Dict[datetime, float]]
    forecast_records: List[ForecastRecord]


def _as_date(v) -> datetime:
    if isinstance(v, datetime):
        return datetime(v.year, v.month, v.day)
    s = str(v).strip().replace('/', '-')
    parts = s.split('-')
    if len(parts) == 3:
        y, m, d = (int(x) for x in parts)
        return datetime(y, m, d)
    return datetime.fromisoformat(s)


def _time_label(v) -> str:
    if isinstance(v, str):
        return v
    if hasattr(v, 'strftime'):
        return v.strftime('%H:%M')
    return str(v)


def _clock_hour(v) -> int:
    if hasattr(v, 'hour'):
        return int(v.hour)
    return int(str(v).strip().split(':')[0])


def plan_slot_start(day: datetime, j: int) -> datetime:
    if not 0 <= int(j) < 144:
        raise IndexError(j)
    return datetime(day.year, day.month, day.day) + timedelta(minutes=10 * (int(j) + 1))


def calendar_slot_start(day: datetime, i: int) -> datetime:
    if not 0 <= int(i) < 144:
        raise IndexError(i)
    return datetime(day.year, day.month, day.day) + timedelta(minutes=10 * int(i))


def plan_slot_index_for_time(day: datetime, t: datetime) -> Optional[int]:
    base = datetime(day.year, day.month, day.day)
    minutes = int(round((t - base).total_seconds() / 60.0))
    j = minutes // 10 - 1
    return int(j) if 0 <= j < 144 else None


def natural_to_plan_ref(day_idx: int, natural_i: int):
    """Return (plan_day_idx, plan_j) for a natural-day interval."""
    if not 0 <= natural_i < 144:
        raise IndexError(natural_i)
    if natural_i == 0:
        return day_idx - 1, 143
    return day_idx, natural_i - 1


def plan_to_natural_ref(plan_day_idx: int, plan_j: int):
    if not 0 <= plan_j < 144:
        raise IndexError(plan_j)
    if plan_j == 143:
        return plan_day_idx + 1, 0
    return plan_day_idx, plan_j + 1


def natural_interval_label(i: int) -> str:
    a = int(i) * 10
    b = (int(i) + 1) * 10
    def fmt(m: int, end: bool = False) -> str:
        if m == 1440:
            return '0:00+1' if end else '24:00'
        return f'{m // 60}:{m % 60:02d}'
    return f'{fmt(a)}-{fmt(b, True)}'


def _plan_to_calendar(arr: np.ndarray) -> np.ndarray:
    arr = np.asarray(arr, dtype=float)
    if arr.shape != (365, 144):
        raise ValueError(f'expected (365,144), got {arr.shape}')
    out = np.full((365, 144), np.nan, dtype=float)
    out[1:, 0] = arr[:-1, 143]
    out[:, 1:] = arr[:, :143]
    return out


def load_q4_inputs(repo_root: Path) -> Q4Data:
    att = Path(repo_root) / '26C题' / '附件'

    wb1 = openpyxl.load_workbook(att / '附件1.xlsx', read_only=True, data_only=True)
    ws1 = wb1[wb1.sheetnames[0]]
    r1 = list(ws1.iter_rows(min_row=2, values_only=True))
    wb1.close()
    if len(r1) != 144:
        raise ValueError('附件1必须是144个固定价时段')
    fixed_price_plan = np.asarray([float(r[1]) for r in r1], dtype=float)
    fixed_matrix = np.tile(fixed_price_plan[None, :], (365, 1))

    wb2 = openpyxl.load_workbook(att / '附件2.xlsx', read_only=True, data_only=True)
    lr = list(wb2['小区负载'].iter_rows(min_row=2, max_row=366, values_only=True))
    pr = list(wb2['光伏发电实际功率'].iter_rows(min_row=2, max_row=366, values_only=True))
    wb2.close()
    if len(lr) != 365 or len(pr) != 365:
        raise ValueError('附件2必须是365天')
    dates: List[datetime] = []
    load_kw = np.zeros((365, 144), dtype=float)
    pv_kw = np.zeros((365, 144), dtype=float)
    for d, (a, b) in enumerate(zip(lr, pr)):
        dates.append(_as_date(a[0]))
        load_kw[d] = np.asarray([float(x or 0.0) for x in a[1:145]], dtype=float)
        pv_kw[d] = np.asarray([float(x or 0.0) for x in b[1:145]], dtype=float)
    if dates[0].date().isoformat() != '2025-01-01' or dates[-1].date().isoformat() != '2025-12-31':
        raise ValueError('附件2日期范围异常')
    load_plan = load_kw * DT
    pv_plan = pv_kw * DT
    net_plan = load_plan - pv_plan

    wb4 = openpyxl.load_workbook(att / '附件4.xlsx', read_only=True, data_only=True)
    ws4 = wb4[wb4.sheetnames[0]]
    r4 = list(ws4.iter_rows(min_row=2, max_row=366, values_only=True))
    if len(r4) != 365 or ws4.max_column != 145:
        wb4.close()
        raise ValueError('附件4必须是365天×144价格点')
    price_dates = [_as_date(r[0]) for r in r4]
    price_plan = np.asarray([[float(x) for x in r[1:145]] for r in r4], dtype=float)
    wb4.close()
    if [x.date() for x in price_dates] != [x.date() for x in dates]:
        raise ValueError('附件2与附件4日期不一致')

    wbt = openpyxl.load_workbook(att / '附件5' / 'result4-2.xlsx', read_only=True, data_only=False)
    ws = wbt['计划购电量']
    plan_headers = [str(ws.cell(1, c).value) for c in range(2, 146)]
    wbt.close()
    if len(plan_headers) != 144 or plan_headers[0] != '0:10-0:20' or plan_headers[-1] != '0:00-0:10+1':
        raise ValueError('result4模板计划时段边界异常')

    forecasts: Dict[datetime, Dict[datetime, float]] = {}
    records: List[ForecastRecord] = []
    wb3 = openpyxl.load_workbook(att / '附件3.xlsx', read_only=True, data_only=True)
    ws3 = wb3[wb3.sheetnames[0]]
    current_date: Optional[datetime] = None
    for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row, values_only=True):
        if row[0] not in (None, ''):
            current_date = _as_date(row[0])
        if current_date is None:
            wb3.close()
            raise ValueError('附件3日期块缺失')
        hour = _clock_hour(row[1])
        issue = current_date + timedelta(hours=hour)
        mp: Dict[datetime, float] = {}
        for lead in range(1, 25):
            target = issue + timedelta(hours=lead)
            val = float(row[1 + lead] or 0.0)
            mp[target] = val
            records.append(ForecastRecord(issue, target, lead, val))
        if issue in forecasts:
            wb3.close()
            raise ValueError(f'附件3重复发行时刻 {issue}')
        forecasts[issue] = mp
    wb3.close()
    if len(forecasts) != 1460 or len(records) != 35040:
        raise ValueError(f'附件3结构异常 issues={len(forecasts)} records={len(records)}')

    return Q4Data(
        dates=dates,
        plan_headers=plan_headers,
        load_plan_kwh=load_plan,
        pv_plan_kwh=pv_plan,
        net_plan_kwh=net_plan,
        load_cal_kwh=_plan_to_calendar(load_plan),
        pv_cal_kwh=_plan_to_calendar(pv_plan),
        net_cal_kwh=_plan_to_calendar(net_plan),
        price_plan=price_plan,
        price_cal=_plan_to_calendar(price_plan),
        fixed_price_plan=fixed_price_plan,
        fixed_price_cal=_plan_to_calendar(fixed_matrix),
        forecasts=forecasts,
        forecast_records=records,
    )


def actual_at(data: Q4Data, t: datetime, field: str) -> Optional[float]:
    d = (t.date() - data.dates[0].date()).days
    i = (t.hour * 60 + t.minute) // 10
    if not (0 <= d < 365 and 0 <= i < 144):
        return None
    arr = {
        'load': data.load_cal_kwh,
        'pv': data.pv_cal_kwh,
        'net': data.net_cal_kwh,
        'price': data.price_cal,
        'fixed_price': data.fixed_price_cal,
    }[field]
    v = float(arr[d, i])
    return v if np.isfinite(v) else None


def grouped_emergency_periods(e: Sequence[float], tol: float = 1e-8):
    e = np.asarray(e, dtype=float)
    idx = np.where(e > tol)[0]
    if idx.size == 0:
        return []
    groups = []
    a = z = int(idx[0])
    for x0 in idx[1:]:
        x = int(x0)
        if x == z + 1:
            z = x
        else:
            groups.append((a, z)); a = z = x
    groups.append((a, z))
    return [(f'{natural_interval_label(a).split("-")[0]}-{natural_interval_label(z).split("-")[1]}', float(e[a:z+1].sum())) for a, z in groups]
