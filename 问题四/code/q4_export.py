# -*- coding: utf-8 -*-
"""Export authoritative Q4 ledgers to the official result4-2/result4-3 templates.

Template geometry/header text is preserved. No formulas, merges, or hidden
postprocessing are introduced.  All numbers are copied from frozen formal
ledgers; the +1 tail remains in each plan row's j=143 contract column.
"""
from __future__ import annotations
from copy import copy
from datetime import datetime,time
from pathlib import Path
import hashlib,json,math
import numpy as np,pandas as pd,openpyxl
from openpyxl.styles import PatternFill,Font,Alignment,Border

ROOT=Path(__file__).resolve().parents[2]; CODE=Path(__file__).resolve().parent
TEMPL=ROOT/'26C题'/'附件'/'附件5'; OUT=ROOT/'问题四'; TABLES=OUT/'output'/'04_tables'; TABLES.mkdir(parents=True,exist_ok=True)
SPEC=['2025-03-20','2025-06-21','2025-09-23','2025-12-21']
POINTS=[(10,0),(12,0),(14,0),(16,0),(18,0),(20,0)]
BLOCKS=[(0,4),(4,8),(8,12),(12,16),(16,20),(20,24)]

def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def _copy_cell_style(dst,src):
    if src.has_style:
        dst._style=copy(src._style); dst.number_format=src.number_format
    if src.alignment: dst.alignment=copy(src.alignment)
    if src.protection: dst.protection=copy(src.protection)

def _fill_plan(ws,ctr,field):
    for ridx,d in enumerate(range(31,365),start=2):
        g=ctr[ctr.day_index==d].sort_values('plan_j')
        if len(g)!=144:raise AssertionError((ws.title,d,len(g)))
        vals=g['B_kwh'].to_numpy(float) if field=='B' else g['A_kwh'].to_numpy(float)
        for j,v in enumerate(vals,start=2): ws.cell(ridx,j,float(v))
        ws.cell(ridx,146,float(vals.sum()))
        fee=float(g['plan_reference_fee_yuan'].sum()) if field=='B' else float(g['final_regular_fee_yuan'].sum())
        ws.cell(ridx,147,fee)
    for row in ws.iter_rows(min_row=2,max_row=335,min_col=2,max_col=147):
        for c in row:
            c.number_format='0.000000'

def _replace_dynamic_rows(ws,records,style_cycle_rows):
    srcstyles=[]
    for rr in style_cycle_rows:
        srcstyles.append([copy(ws.cell(rr,c)._style) for c in range(1,ws.max_column+1)])
    ws.delete_rows(2,ws.max_row-1)
    for i,rec in enumerate(records,start=2):
        ws.insert_rows(i)
        sty=srcstyles[(i-2)%len(srcstyles)]
        for c in range(1,ws.max_column+1): ws.cell(i,c)._style=copy(sty[c-1])
        for c,v in enumerate(rec,start=1): ws.cell(i,c,v)

def _stabilize_display_widths(wb):
    """Presentation-only widths to prevent Excel date cells rendering as ####."""
    if '充放电量' in wb.sheetnames:
        ws=wb['充放电量']
        for col,width in {'A':14,'B':15,'C':15,'D':15,'E':12,'F':15}.items():
            ws.column_dimensions[col].width=max(ws.column_dimensions[col].width or 0,width)
    if '紧急购电量' in wb.sheetnames:
        ws=wb['紧急购电量']
        for col,width in {'A':14,'B':18,'C':15}.items():
            ws.column_dimensions[col].width=max(ws.column_dimensions[col].width or 0,width)

def _charge_records(slots,days):
    rec=[]
    for ds in days.date:
        ss=slots[slots.date==ds].sort_values('slot'); dd=days[days.date==ds].iloc[0]
        if len(ss)!=144:raise AssertionError((ds,len(ss)))
        for bi,(a,b) in enumerate(BLOCKS):
            sub=ss[(ss.slot>=a*6)&(ss.slot<b*6)]
            rec.append([datetime.fromisoformat(ds) if bi==0 else None,f'{a}:00-{b}:00',float(sub.x.sum()),float(sub.y.sum()),
                        time(0,0) if bi==0 else ('24:00' if bi==1 else None),float(dd.S00_kwh) if bi==0 else (float(dd.S24_kwh) if bi==1 else None)])
    return rec

def _emergency_records(slots,days,tol=1e-8):
    rec=[]
    for ds in days.date:
        s=slots[slots.date==ds].sort_values('slot'); e=s.e_kwh.to_numpy(float); idx=np.where(e>tol)[0]
        groups=[]
        if len(idx):
            a=z=int(idx[0])
            for q in idx[1:]:
                q=int(q)
                if q==z+1:z=q
                else:groups.append((a,z));a=z=q
            groups.append((a,z))
        if not groups: rec.append([datetime.fromisoformat(ds),'无',0.0]); continue
        for gi,(a,z) in enumerate(groups):
            sm=a*10; em=(z+1)*10
            def fmt(m): return '24:00' if m==1440 else f'{m//60}:{m%60:02d}'
            rec.append([datetime.fromisoformat(ds) if gi==0 else None,f'{fmt(sm)}-{fmt(em)}',float(e[a:z+1].sum())])
    return rec

def export_one(br,fn):
    cdir=CODE/br; slots=pd.read_csv(cdir/'physical_10min.csv'); ctr=pd.read_csv(cdir/'contract_ledger.csv'); days=pd.read_csv(cdir/'daily_ledger.csv')
    template=TEMPL/fn; out=OUT/fn
    wb=openpyxl.load_workbook(template)
    _fill_plan(wb['计划购电量'],ctr,'B')
    if br=='q4_3': _fill_plan(wb['调整购电量'],ctr,'A')
    _replace_dynamic_rows(wb['充放电量'],_charge_records(slots,days),[2,3,4,5,6,7])
    _replace_dynamic_rows(wb['紧急购电量'],_emergency_records(slots,days),[2,3,4,5])
    _stabilize_display_widths(wb)
    # Normalize dynamic date columns: style cycling in the official example
    # otherwise alternates between short dates and a full datetime display.
    for sheet in ('充放电量','紧急购电量'):
        ws=wb[sheet]
        for cell in ws['A'][1:]:
            if isinstance(cell.value,datetime): cell.number_format='yyyy/m/d'
    # Stable numeric/date formatting; no formulas or merges.
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value,float): cell.number_format='0.000000'
    wb.calculation.fullCalcOnLoad=False; wb.calculation.forceFullCalc=False
    wb.save(out)
    return {'path':str(out),'sha256':sha(out),'sheets':wb.sheetnames}

def specified_tables(br):
    cdir=CODE/br; slots=pd.read_csv(cdir/'physical_10min.csv'); days=pd.read_csv(cdir/'daily_ledger.csv')
    rows1=[];rows2=[];rows3=[]
    for ds in SPEC:
        s=slots[slots.date==ds].sort_values('slot'); d=days[days.date==ds].iloc[0]
        if len(s)!=144:raise AssertionError(ds)
        for h,m in POINTS:
            i=(h*60+m)//10; r=s[s.slot==i].iloc[0]
            rows1.append({'branch':br,'date':ds,'interval':f'{h:02d}:00-{h:02d}:10','contract_purchase_kwh':float(r.A_kwh),
                          'realized_grid_import_kwh':float(r.I_kwh),'price_yuan_per_kwh':float(r.price_yuan_per_kwh),
                          'all_day_contract_kwh':float(s.A_kwh.sum()),'all_day_cash_yuan':float(d.cash_fee_yuan)})
        for a,b in BLOCKS:
            q=s[(s.slot>=a*6)&(s.slot<b*6)]
            rows2.append({'branch':br,'date':ds,'block':f'{a:02d}:00-{b:02d}:00','charge_kwh':float(q.x.sum()),'discharge_kwh':float(q.y.sum()),
                          'S00_kwh':float(d.S00_kwh),'S24_kwh':float(d.S24_kwh)})
        e=s.e_kwh.to_numpy(float); idx=np.where(e>1e-8)[0]
        if len(idx)==0:rows3.append({'branch':br,'date':ds,'interval':'无','emergency_kwh':0.0})
        else:
            a=z=int(idx[0]);groups=[]
            for q0 in idx[1:]:
                q0=int(q0)
                if q0==z+1:z=q0
                else:groups.append((a,z));a=z=q0
            groups.append((a,z))
            for a,z in groups:
                def fmt(i,end=False):
                    m=(i+(1 if end else 0))*10
                    return '24:00' if m==1440 else f'{m//60:02d}:{m%60:02d}'
                rows3.append({'branch':br,'date':ds,'interval':f'{fmt(a)}-{fmt(z,True)}','emergency_kwh':float(e[a:z+1].sum())})
    return pd.DataFrame(rows1),pd.DataFrame(rows2),pd.DataFrame(rows3)

def run():
    manifest={'generator':'q4_export.py','spreadsheet_engine':'openpyxl template-preserving fallback because artifact_tool is unavailable inside the project Python runtime','files':{}}
    manifest['files']['result4-2']=export_one('q4_2','result4-2.xlsx'); manifest['files']['result4-3']=export_one('q4_3','result4-3.xlsx')
    all1=[];all2=[];all3=[]
    for br in ('q4_2','q4_3'):
        a,b,c=specified_tables(br);all1.append(a);all2.append(b);all3.append(c)
        a.to_csv(TABLES/f'{br}_specified_purchase.csv',index=False,encoding='utf-8-sig');b.to_csv(TABLES/f'{br}_specified_storage.csv',index=False,encoding='utf-8-sig');c.to_csv(TABLES/f'{br}_specified_emergency.csv',index=False,encoding='utf-8-sig')
    A=pd.concat(all1,ignore_index=True);B=pd.concat(all2,ignore_index=True);C=pd.concat(all3,ignore_index=True)
    A.to_csv(TABLES/'specified_purchase_all.csv',index=False,encoding='utf-8-sig');B.to_csv(TABLES/'specified_storage_all.csv',index=False,encoding='utf-8-sig');C.to_csv(TABLES/'specified_emergency_all.csv',index=False,encoding='utf-8-sig')
    lines=['# 问题四指定日期权威表','', '> 六个10分钟点严格按题面表1：10:00、12:00、14:00、16:00、18:00、20:00。购电量主列为交付时最终合同量A；同时CSV保留实际使用I。全天购电费为权威现金账（正常合同结算+紧急购电）。','']
    for br in ('q4_2','q4_3'):
        lines+=['## '+br,'','### 表1 六个10分钟购电点']
        x=A[A.branch==br]
        for ds in SPEC:
            y=x[x.date==ds];lines+=['',f'**{ds}**', '', '|时段|最终合同购电量/kWh|实际外网使用/kWh|电价/(元/kWh)|全天合同量/kWh|全天现金费用/元|','|---|---:|---:|---:|---:|---:|']
            for _,r in y.iterrows():lines.append(f"|{r.interval}|{r.contract_purchase_kwh:.3f}|{r.realized_grid_import_kwh:.3f}|{r.price_yuan_per_kwh:.4f}|{r.all_day_contract_kwh:.3f}|{r.all_day_cash_yuan:.2f}|")
        lines+=['','### 表2 六个四小时块充放电及S00/S24','']
        for ds in SPEC:
            y=B[(B.branch==br)&(B.date==ds)];lines += [f'**{ds}**','', '|时段|充电量/kWh|放电量/kWh|S00/kWh|S24/kWh|','|---|---:|---:|---:|---:|']
            for _,r in y.iterrows():lines.append(f"|{r.block}|{r.charge_kwh:.3f}|{r.discharge_kwh:.3f}|{r.S00_kwh:.3f}|{r.S24_kwh:.3f}|")
            lines.append('')
        lines+=['### 表3 紧急购电事件','', '|日期|时段|紧急购电量/kWh|','|---|---|---:|']
        for _,r in C[C.branch==br].iterrows():lines.append(f"|{r.date}|{r.interval}|{r.emergency_kwh:.3f}|")
        lines.append('')
    (TABLES/'指定日期表格.md').write_text('\n'.join(lines),encoding='utf-8')
    manifest['specified_table_hashes']={p.name:sha(p) for p in TABLES.iterdir() if p.is_file()}
    (CODE/'export_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(manifest,ensure_ascii=False),flush=True)
if __name__=='__main__':run()
