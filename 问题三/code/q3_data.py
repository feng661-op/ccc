# -*- coding: utf-8 -*-
from __future__ import annotations
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple
import numpy as np
import openpyxl

DT=1/6; CAP=12000.0; PMAX=5000.0; XMAX=PMAX*DT
SOC0=6000.0; SOC_MIN=1200.0; SOC_MAX=10800.0
ETA_C=float(np.sqrt(0.9)); ETA_D=float(np.sqrt(0.9))
EVENT_HOURS=(0,6,12,18); EVENT_PLAN_START={0:0,6:35,12:71,18:107}; EVENT_PLAN_HORIZON={0:145,6:109,12:73,18:37}; EVAL_START=31

@dataclass(frozen=True)
class ForecastRecord:
    issue_time: datetime; target_time: datetime; lead_hour: int; pv_kw: float

@dataclass
class Q3Data:
    dates: List[datetime]
    price_plan: np.ndarray; price_calendar: np.ndarray
    load_plan_kwh: np.ndarray; pv_plan_kwh: np.ndarray; net_plan_kwh: np.ndarray
    load_cal_kwh: np.ndarray; pv_cal_kwh: np.ndarray; net_cal_kwh: np.ndarray
    plan_headers: List[str]
    forecasts: Dict[datetime, Dict[datetime,float]]
    forecast_records: List[ForecastRecord]

def plan_slot_start(day:datetime,j:int)->datetime:
    return datetime(day.year,day.month,day.day)+timedelta(minutes=10*(j+1))

def calendar_slot_start(day:datetime,i:int)->datetime:
    return datetime(day.year,day.month,day.day)+timedelta(minutes=10*i)

def plan_slot_index_for_time(day:datetime,t:datetime)->Optional[int]:
    base=datetime(day.year,day.month,day.day)
    m=int(round((t-base).total_seconds()/60)); j=m//10-1
    return j if 0<=j<144 else None

def price_for_time(data:Q3Data,t:datetime)->float:
    m=t.hour*60+t.minute
    return float(data.price_plan[-1] if m<10 else data.price_plan[m//10-1])

def _clock_hour(v)->int:
    if hasattr(v,'hour'): return int(v.hour)
    return int(str(v).strip().split(':')[0])

def load_q3_inputs(repo_root:Path)->Q3Data:
    att=repo_root/'26C题'/'附件'
    wb=openpyxl.load_workbook(att/'附件1.xlsx',read_only=True,data_only=True); ws=wb['Sheet1']
    r1=list(ws.iter_rows(min_row=2,values_only=True)); wb.close()
    if len(r1)!=144: raise ValueError('附件1非144点')
    pp=np.asarray([float(r[1]) for r in r1]); pc=np.concatenate(([pp[-1]],pp[:143]))
    wb=openpyxl.load_workbook(att/'附件2.xlsx',read_only=True,data_only=True)
    lr=list(wb['小区负载'].iter_rows(min_row=2,max_row=366,values_only=True)); pr=list(wb['光伏发电实际功率'].iter_rows(min_row=2,max_row=366,values_only=True)); wb.close()
    if len(lr)!=365 or len(pr)!=365: raise ValueError('附件2非365天')
    dates=[]; lk=np.zeros((365,144)); pk=np.zeros((365,144))
    for d,(a,b) in enumerate(zip(lr,pr)):
        dv=a[0]
        if isinstance(dv,datetime): dates.append(datetime(dv.year,dv.month,dv.day))
        else:
            y,m,dd=[int(x) for x in str(dv).strip().replace('/','-').split('-')]; dates.append(datetime(y,m,dd))
        lk[d]=[float(x or 0) for x in a[1:145]]; pk[d]=[float(x or 0) for x in b[1:145]]
    lp=lk*DT; pv=pk*DT; net=lp-pv
    lc=np.full((365,144),np.nan); gc=np.full((365,144),np.nan)
    for d in range(365):
        if d>=1: lc[d,0]=lp[d-1,143]; gc[d,0]=pv[d-1,143]
        lc[d,1:]=lp[d,:143]; gc[d,1:]=pv[d,:143]
    nc=lc-gc
    wb=openpyxl.load_workbook(att/'附件5'/'result3.xlsx',read_only=True,data_only=True); ws=wb['计划购电量']
    headers=[str(ws.cell(1,c).value) for c in range(2,146)]; wb.close()
    if headers[0]!='0:10-0:20' or headers[-1]!='0:00-0:10+1': raise ValueError('result3时间边界异常')
    wb=openpyxl.load_workbook(att/'附件3.xlsx',read_only=True,data_only=True); ws=wb['Sheet1']
    forecasts={}; records=[]; current=None
    for row in ws.iter_rows(min_row=2,max_row=ws.max_row,values_only=True):
        if row[0] not in (None,''):
            y,m,dd=[int(x) for x in str(row[0]).strip().replace('/','-').split('-')]; current=datetime(y,m,dd)
        if current is None: raise ValueError('附件3日期为空')
        h=_clock_hour(row[1]); issue=current+timedelta(hours=h); mp={}
        for lead in range(1,25):
            target=issue+timedelta(hours=lead); val=float(row[1+lead] or 0); mp[target]=val; records.append(ForecastRecord(issue,target,lead,val))
        if issue in forecasts: raise ValueError('附件3重复issue')
        forecasts[issue]=mp
    wb.close()
    if len(forecasts)!=1460 or len(records)!=35040: raise ValueError('附件3结构异常')
    return Q3Data(dates,pp,pc,lp,pv,net,lc,gc,nc,headers,forecasts,records)

def completed_actual_kw(data:Q3Data,t:datetime,kind:str)->Optional[float]:
    start=t-timedelta(minutes=10); d=(start.date()-data.dates[0].date()).days
    if not 0<=d<365: return None
    i=(start.hour*60+start.minute)//10; arr=data.pv_cal_kwh if kind=='pv' else data.load_cal_kwh; v=arr[d,i]
    return None if not np.isfinite(v) else float(v/DT)

def historical_profile_kw(data:Q3Data,issue:datetime,kind:str,horizon:int)->np.ndarray:
    """Strictly causal historical forecast.

    Load uses a 7-day seasonal-naive baseline whenever the target's same slot
    one week earlier is available. This choice is fixed using January only; the
    formal-period comparison is reporting, not tuning. PV keeps the robust
    historical fallback used outside the official-vintage support.
    """
    arr=data.pv_cal_kwh if kind=='pv' else data.load_cal_kwh; dcur=(issue.date()-data.dates[0].date()).days; out=np.zeros(horizon)
    for k in range(horizon):
        t=issue+timedelta(minutes=10*k); si=(t.hour*60+t.minute)//10; target_d=(t.date()-data.dates[0].date()).days
        if kind=='load':
            j=target_d-7
            if 0<=j<365 and j<dcur and np.isfinite(arr[j,si]):
                out[k]=max(0,float(arr[j,si])/DT); continue
        vals=[]; same=[]
        for lag in range(1,8):
            j=dcur-lag
            if 0<=j<365 and np.isfinite(arr[j,si]): vals.append(float(arr[j,si]))
        for j in range(max(0,dcur-35),dcur):
            if data.dates[j].weekday()==t.weekday() and np.isfinite(arr[j,si]): same.append(float(arr[j,si]))
        base=float(np.median(vals)) if vals else (float(np.median(same)) if same else (0 if kind=='pv' else 5000*DT))
        if same: base=.75*base+.25*float(np.median(same[-4:]))
        out[k]=max(0,base/DT)
    return out

def _selected_vintage_hour(event_hour:int,use_new_vintage:bool,disabled_vintage_hours:Sequence[int]=())->int:
    if not use_new_vintage: return 0
    disabled=set(int(x) for x in disabled_vintage_hours)
    if event_hour not in disabled: return event_hour
    return max(h for h in EVENT_HOURS if h<event_hour and h not in disabled) if event_hour>0 else 0

def _previous_used_vintage_hour(event_hour:int,disabled_vintage_hours:Sequence[int]=())->int:
    disabled=set(int(x) for x in disabled_vintage_hours)
    prev=[h for h in EVENT_HOURS if h<event_hour and h not in disabled]
    return max(prev) if prev else 0

def official_pv_kw(data:Q3Data,current_issue:datetime,vintage_issue:datetime,times:Sequence[datetime])->np.ndarray:
    fallback=historical_profile_kw(data,current_issue,'pv',len(times)); mp=data.forecasts.get(vintage_issue)
    if not mp: return fallback
    ts=sorted(mp); xs=np.asarray([(x-vintage_issue).total_seconds()/3600 for x in ts]); ys=np.asarray([mp[x] for x in ts])
    anchor=completed_actual_kw(data,vintage_issue,'pv') or 0.0; xs=np.r_[0.0,xs]; ys=np.r_[anchor,ys]; out=np.zeros(len(times))
    for k,t in enumerate(times):
        rel=(t-vintage_issue).total_seconds()/3600
        out[k]=np.interp(rel,xs,ys) if 0<=rel<=24 else fallback[k]
    return np.maximum(out,0)

def event_point_forecast(data:Q3Data,day_idx:int,event_hour:int,horizon:int,use_new_vintage:bool=True,disabled_vintage_hours:Sequence[int]=())->Tuple[np.ndarray,List[datetime],np.ndarray]:
    day=data.dates[day_idx]; issue=datetime(day.year,day.month,day.day)+timedelta(hours=event_hour)
    times=[issue+timedelta(minutes=10*k) for k in range(horizon)]
    load=historical_profile_kw(data,issue,'load',horizon)
    vh=_selected_vintage_hour(event_hour,use_new_vintage,disabled_vintage_hours)
    vint=datetime(day.year,day.month,day.day)+timedelta(hours=vh)
    pv=official_pv_kw(data,issue,vint,times)
    if event_hour>0:
        rs=[]
        for back in range(1,7):
            t=issue-timedelta(minutes=10*back); d=(t.date()-data.dates[0].date()).days; i=(t.hour*60+t.minute)//10
            if 0<=d<365 and np.isfinite(data.load_cal_kwh[d,i]):
                actual=float(data.load_cal_kwh[d,i]/DT); pred=historical_profile_kw(data,t,'load',1)[0]; rs.append(actual-pred)
        if rs: load=np.maximum(0,load+float(np.median(rs))*np.exp(-np.arange(horizon)/36.0))
    net=(load-pv)*DT; prices=np.asarray([price_for_time(data,t) for t in times]); return net,times,prices

def actual_horizon(data:Q3Data,issue:datetime,horizon:int)->Optional[np.ndarray]:
    out=np.zeros(horizon)
    for k in range(horizon):
        t=issue+timedelta(minutes=10*k); d=(t.date()-data.dates[0].date()).days; i=(t.hour*60+t.minute)//10
        if not 0<=d<365 or not np.isfinite(data.net_cal_kwh[d,i]): return None
        out[k]=float(data.net_cal_kwh[d,i])
    return out

def prefix_signal(data:Q3Data,day_idx:int,event_hour:int,use_new_vintage:bool,disabled_vintage_hours:Sequence[int]=())->np.ndarray:
    day=data.dates[day_idx]; base=datetime(day.year,day.month,day.day); vals=[]
    if event_hour>0:
        issue=base+timedelta(hours=event_hour)
        for back in range(1,7):
            t=issue-timedelta(minutes=10*back); d=(t.date()-data.dates[0].date()).days; i=(t.hour*60+t.minute)//10
            if 0<=d<365 and np.isfinite(data.net_cal_kwh[d,i]): vals.append(float(data.net_cal_kwh[d,i]))
    out=[float(np.mean(vals)) if vals else 0.0]
    disabled=set(int(x) for x in disabled_vintage_hours)
    if use_new_vintage and event_hour>0 and event_hour not in disabled:
        oldh=_previous_used_vintage_hour(event_hour,disabled); old=base+timedelta(hours=oldh); new=base+timedelta(hours=event_hour); a=data.forecasts.get(old,{}); b=data.forecasts.get(new,{})
        shared=[t for t in b if t in a and t>new]
        if shared:
            inn=np.asarray([b[t]-a[t] for t in shared]); out += [float(np.mean(inn)),float(np.mean(np.abs(inn)))]
        else: out += [0.0,0.0]
    else: out += [0.0,0.0]
    return np.asarray(out)

def forecast_alignment_rows(data:Q3Data):
    issues=sorted(data.forecasts); out=[]
    for old,new in zip(issues[:-1],issues[1:]):
        if (new-old).total_seconds()!=6*3600: continue
        a,b=data.forecasts[old],data.forecasts[new]
        for target in sorted(set(a)&set(b)):
            lo=int(round((target-old).total_seconds()/3600)); ln=int(round((target-new).total_seconds()/3600)); unreal=target>new
            out.append((old,new,target,lo,ln,unreal,unreal))
    return out

def settlement_components(b,a,p_delivery,down_settlement='cancel_settlement',adjustment_price=None):
    b=np.asarray(b,float); a=np.asarray(a,float); p=np.asarray(p_delivery,float); pa=p if adjustment_price is None else np.asarray(adjustment_price,float)
    down=np.maximum(b-a,0); up=np.maximum(a-b,0); keep=np.minimum(a,b)
    if down_settlement=='cancel_settlement': plan=p*keep
    elif down_settlement=='sunk_plan_plus_penalty': plan=p*b
    else: raise ValueError(down_settlement)
    cancel=.5*pa*down; add=1.5*pa*up
    return {'F_plan':plan,'F_cancel':cancel,'F_add':add,'F_regular':plan+cancel+add,'down_kwh':down,'up_kwh':up,'keep_kwh':keep}

def grouped_emergency_periods(e,tol=1e-8):
    e=np.asarray(e,float); idx=np.where(e>tol)[0]
    if idx.size==0:return []
    groups=[]; a=z=int(idx[0])
    for x0 in idx[1:]:
        x=int(x0)
        if x==z+1:z=x
        else:groups.append((a,z));a=z=x
    groups.append((a,z)); out=[]
    for a,z in groups:
        sm=a*10; em=(z+1)*10
        def fmt(m,end=False):
            if m==1440:return '0:00+1' if end else '24:00'
            return f'{m//60}:{m%60:02d}'
        out.append((f'{fmt(sm)}-{fmt(em,True)}',float(e[a:z+1].sum())))
    return out
