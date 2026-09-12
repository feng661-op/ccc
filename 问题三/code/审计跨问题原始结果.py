"""Read-only audit of existing Q2/Q3, with evidence confined to Q3."""
from pathlib import Path
import csv
import hashlib
import json
import sys
import time
import zipfile
import numpy as np
import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
OUT = HERE / 'cross_problem_evidence'
OUT.mkdir(exist_ok=True)
sys.path.insert(0, str(HERE))
from q3_data import load_q3_inputs, settlement_components
from q3_sim import simulate_strategy


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def main():
    if (OUT/'before_comparison.json').exists():
        raise SystemExit('Original audit is frozen. Read before_audit.zip and old_closed_replay.json; refusing to replace original evidence with corrected results.')
    # Preserve every original Q3 file, including prior physical-fix evidence.
    backup = OUT / 'before_audit.zip'
    if not backup.exists():
        with zipfile.ZipFile(backup, 'w', zipfile.ZIP_DEFLATED) as archive:
            for p in sorted(HERE.parent.rglob('*')):
                if p.is_file() and OUT not in p.parents and not any(x in p.parts for x in ('__pycache__', '.pytest_cache')):
                    archive.write(p, p.relative_to(ROOT))
    protected = {str(p.relative_to(ROOT)): sha(p) for folder in ('问题一', '问题二', '问题四', '26C题')
                 for p in (ROOT/folder).rglob('*') if p.is_file() and '__pycache__' not in p.parts and '.pytest_cache' not in p.parts}
    (OUT/'protected_hashes.json').write_text(json.dumps(protected, ensure_ascii=False, indent=2), encoding='utf-8')
    inventory = []
    for folder in ('问题二', '问题三'):
        for p in sorted((ROOT/folder).rglob('*')):
            if not p.is_file() or OUT in p.parents or any(x in p.parts for x in ('__pycache__', '.pytest_cache')):
                continue
            item = {'path': str(p.relative_to(ROOT)), 'sha256': sha(p), 'bytes': p.stat().st_size}
            if p.suffix == '.json':
                value = json.loads(p.read_text(encoding='utf-8-sig'))
                item['top_keys'] = list(value) if isinstance(value, dict) else len(value)
            elif p.suffix == '.csv':
                with p.open(encoding='utf-8-sig', newline='') as f:
                    reader = csv.reader(f); item['columns'] = next(reader, [])
                    item['rows'] = sum(1 for _ in reader)
            elif p.suffix == '.npz':
                with np.load(p) as z:
                    item['arrays'] = {k: {'shape': list(z[k].shape), 'finite': bool(np.isfinite(z[k]).all()) if z[k].dtype.kind in 'fiu' else None} for k in z.files}
            elif p.suffix == '.xlsx':
                wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
                item['sheets'] = {ws.title: {'shape': [ws.max_row, ws.max_column], 'nonempty': sum(v is not None for row in ws.iter_rows(values_only=True) for v in row)} for ws in wb}
                wb.close()
            inventory.append(item)
    (OUT/'original_inventory.json').write_text(json.dumps(inventory, ensure_ascii=False, indent=2), encoding='utf-8')
    data = load_q3_inputs(ROOT)
    q2 = np.load(ROOT/'问题二'/'code'/'run_data.npz')
    e = slice(31, None)
    q = np.zeros_like(q2['plan_q']); q[:, 1:] = q2['plan_q'][:, :143]; q[1:, 0] = q2['plan_q'][:-1, 143]
    x, y, emergency = (q2[k] for k in ('charge', 'discharge', 'emergency'))
    need = data.load_cal_kwh + x - y
    cur = np.maximum(data.pv_cal_kwh-need, 0)
    grid = np.minimum(q, np.maximum(need-data.pv_cal_kwh, 0))
    unused = q-grid
    sp = np.column_stack((q2['soc00'], q2['soc00'][:, None] + np.cumsum(np.sqrt(.9)*x-y/np.sqrt(.9), axis=1)))
    bridge = float(data.price_plan[-1]*(q2['plan_q'][30, -1]-q2['plan_q'][-1, -1]))
    out = {'Q2': dict(plan_kwh=float(q2['plan_q'][e].sum()), accepted_plan_kwh=float(q2['plan_q'][e].sum()),
        natural_contract_kwh=float(q[e].sum()), actual_regular_import_kwh=float(grid[e].sum()), unused_contract_kwh=float(unused[e].sum()),
        charge_kwh=float(x[e].sum()), discharge_kwh=float(y[e].sum()), emergency_kwh=float(emergency[e].sum()),
        pv_curtail_kwh=float(cur[e].sum()), total_surplus_kwh=float(q2['curtailment'][e].sum()),
        soc_min=float(sp[e].min()), soc_max=float(sp[e].max()), soc_start=float(sp[31, 0]), soc_end=float(sp[-1, -1]),
        plan_cost=float(q2['plan_cost'][e].sum()), emergency_cost=float(q2['emergency_cost'][e].sum()),
        published_total=float((q2['plan_cost'][e]+q2['emergency_cost'][e]).sum()), natural_bridge=bridge,
        natural_total=float((q2['plan_cost'][e]+q2['emergency_cost'][e]).sum()+bridge),
        surplus_split_error=float(np.max(np.abs(cur[e]+unused[e]-q2['curtailment'][e]))),
        physical_error=float(np.max(np.abs(grid[e]+data.pv_cal_kwh[e]-cur[e]+y[e]+emergency[e]-data.load_cal_kwh[e]-x[e]))))}
    for name in ('A', 'B', 'C', 'D'):
        z = np.load(HERE/f'run_{name}.npz'); comp = settlement_components(z['B'][e], z['A'][e], data.price_plan)
        out['old_'+name] = {k:float(z[k][e].sum()) for k in ('charge','discharge','emergency','curtail','grid_import','unused_contract','supply_surplus','natural_regular_fee','emergency_fee')}
        out['old_'+name].update(plan_kwh=float(z['B'][e].sum()), final_kwh=float(z['A'][e].sum()),
            up_kwh=float(comp['up_kwh'].sum()), down_kwh=float(comp['down_kwh'].sum()),
            soc_start=float(z['soc00'][31]), soc_end=float(z['soc24'][-1]),
            natural_total=float((z['natural_regular_fee'][e]+z['emergency_fee'][e]).sum()))
    print(json.dumps(out, ensure_ascii=False, indent=2), flush=True)
    (OUT/'before_comparison.json').write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding='utf-8')
    start = time.time()
    s = simulate_strategy(data, name='old_closed', use_new_vintage=False, allow_revision=False, scenario_count=3, verbose=True)
    old = np.load(HERE/'run_A.npz')
    replay = {'seconds':time.time()-start, 'total_cost':float((s.natural_regular_fee[e]+s.emergency_fee[e]).sum()),
              'array_max_errors':{k:float(np.max(np.abs(getattr(s,k)-old[k]))) for k in ('B','A','charge','discharge','emergency','soc_path')}}
    (OUT/'old_closed_replay.json').write_text(json.dumps(replay, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(replay, ensure_ascii=False, indent=2), flush=True)


if __name__ == '__main__':
    main()
