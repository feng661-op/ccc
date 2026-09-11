# -*- coding: utf-8 -*-
"""
问题1 模型求解
==============
确定性储能套利线性规划(LP)，两阶段字典序求解：
  阶段1：min 购电费  sum(c_t * b_t)
  阶段2：在费用已最优的前提下 min 吞吐量 sum(x_t + y_t)（消除多最优解）

输出：
  - D:/数学建模/26C题/问题一/result1.xlsx  （复制附件5模板后填值）
  - code/metrics.json              （关键指标，供绘图与文档使用）
"""
import os
import json
import shutil
import sys

import numpy as np
from scipy.optimize import linprog

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from 数据读取与预处理 import load_data  # noqa: E402

BASE = r'D:\数学建模\26C题'
CODE = os.path.join(BASE, '问题一', 'code')
TEMPLATE = os.path.join(BASE, '附件', '附件5', 'result1.xlsx')
RESULT = os.path.join(BASE, '问题一', 'result1.xlsx')

# ---- 储能参数(附录1) ----
CAP = 12000.0        # 最大容量 kWh
PMAX = 5000.0        # 最大充放电功率 kW
SOC0 = 6000.0        # 初始/末态电量 kWh
SOC_MIN = 1200.0
SOC_MAX = 10800.0
ETA_C = np.sqrt(0.9)  # 充电效率 = √0.9 ≈ 0.9487（往返90%无偏对称拆分）
ETA_D = np.sqrt(0.9)  # 放电效率 = √0.9 ≈ 0.9487
DT = 1.0 / 6.0       # 10分钟 = 1/6 小时
XMAX = PMAX * DT     # 每段最大充/放能量 = 833.33 kWh


def build_and_solve(data, eta_c=ETA_C, eta_d=ETA_D, soc_min=SOC_MIN,
                    soc_max=SOC_MAX, xmax=XMAX):
    """构建并求解 LP，返回最优解字典。"""
    n = data['n']
    price = np.asarray(data['price'], dtype=float)
    load = np.asarray(data['load'], dtype=float)
    pv = np.asarray(data['pv'], dtype=float)

    # 样本=区间起点：样本时刻 T 对应区间 [T, T+10min]。
    # 段 i=[i*10,(i+1)*10] 使用样本 (i-1) mod n（段0 用时刻24:00样本，日周期24:00≡0:00）。
    load = np.roll(load, 1)
    pv = np.roll(pv, 1)
    price = np.roll(price, 1)

    # 变量布局：v = [b(0..n-1), x(n..2n-1), y(2n..3n-1), S(3n..4n)]
    B = 0
    X = n
    Y = 2 * n
    S = 3 * n
    N = 3 * n + (n + 1)          # 总变量数 = 144*3 + 145 = 577

    # 目标
    c1 = np.zeros(N)             # 阶段1：购电费
    c1[B:B + n] = price
    c2 = np.zeros(N)             # 阶段2：吞吐量(充+放)
    c2[X:X + n] = 1.0
    c2[Y:Y + n] = 1.0

    # 等式：SOC递推 + 首末固定
    Aeq = np.zeros((n + 2, N))
    beq = np.zeros(n + 2)
    for t in range(n):
        Aeq[t, S + t + 1] = 1.0          # S[t+1]
        Aeq[t, S + t] = -1.0             # - S[t]
        Aeq[t, X + t] = -eta_c           # - eta_c * x[t]
        Aeq[t, Y + t] = 1.0 / eta_d      # + y[t]/eta_d
    Aeq[n, S + 0] = 1.0; beq[n] = SOC0        # S[0] = 6000
    Aeq[n + 1, S + n] = 1.0; beq[n + 1] = SOC0  # S[n] = 6000

    # 不等式：供需平衡 b + pv + y >= load + x  <=>  x - y - b <= pv - load
    Aub = np.zeros((n, N))
    bub = np.zeros(n)
    for t in range(n):
        Aub[t, X + t] = 1.0
        Aub[t, Y + t] = -1.0
        Aub[t, B + t] = -1.0
        bub[t] = pv[t] - load[t]

    # 边界
    bounds = []
    for i in range(N):
        if B <= i < B + n:
            bounds.append((0.0, None))            # b >= 0
        elif X <= i < X + n:
            bounds.append((0.0, xmax))            # 0 <= x <= xmax
        elif Y <= i < Y + n:
            bounds.append((0.0, xmax))            # 0 <= y <= xmax
        else:
            bounds.append((soc_min, soc_max))     # soc_min <= S <= soc_max

    # 阶段1：最小购电费
    r1 = linprog(c1, A_ub=Aub, b_ub=bub, A_eq=Aeq, b_eq=beq,
                 bounds=bounds, method='highs')
    if not r1.success:
        raise RuntimeError(f'阶段1求解失败: {r1.message}')
    f_star = float(r1.fun)

    # 阶段2：固定费用最优，最小吞吐量(消多解)
    Aub2 = np.vstack([Aub, c1.reshape(1, -1)])
    bub2 = np.concatenate([bub, [f_star + 1e-5]])
    r2 = linprog(c2, A_ub=Aub2, b_ub=bub2, A_eq=Aeq, b_eq=beq,
                 bounds=bounds, method='highs')
    v = r2.x if r2.success else r1.x

    b = v[B:B + n]
    x = v[X:X + n]
    y = v[Y:Y + n]
    s = v[S:S + n + 1]

    return {
        'cost': float(c1 @ v),                  # 最优购电费(元)
        'b': b, 'x': x, 'y': y, 's': s,
        'total_purchase': float(b.sum()),       # 全天购电量 kWh
        'total_charge': float(x.sum()),         # 全天充电量 kWh
        'total_discharge': float(y.sum()),      # 全天放电量 kWh
        'soc_min': float(s.min()),
        'soc_max': float(s.max()),
    }


def write_result1(sol, data):
    """复制附件5模板，填值保存到 result1.xlsx。"""
    if os.path.exists(RESULT):
        os.remove(RESULT)
    shutil.copy(TEMPLATE, RESULT)

    import openpyxl
    wb = openpyxl.load_workbook(RESULT)

    # 「计划购电量」：模板行 t(0起) 标签为 [10(t+1), 10(t+2)] 分钟，对应段 (t+1) mod n
    ws = wb['计划购电量']
    n = data['n']
    for t in range(n):
        ws.cell(row=2 + t, column=2, value=round(float(sol['b'][(t + 1) % n]), 4))

    # 「充放电量」：6 个 4 小时块
    ws2 = wb['充放电量']
    blocks = [(0, 24), (24, 48), (48, 72), (72, 96), (96, 120), (120, 144)]
    for i, (a, z) in enumerate(blocks):
        ws2.cell(row=2 + i, column=2, value=round(float(sol['x'][a:z].sum()), 4))
        ws2.cell(row=2 + i, column=3, value=round(float(sol['y'][a:z].sum()), 4))
    ws2.cell(row=2, column=5, value=round(float(sol['s'][0]), 4))    # 0:00 储电量
    ws2.cell(row=3, column=5, value=round(float(sol['s'][-1]), 4))   # 24:00 储电量

    wb.save(RESULT)


def main():
    data = load_data()
    sol = build_and_solve(data)
    write_result1(sol, data)

    price = np.asarray(data['price'])
    load = np.asarray(data['load'])
    pv = np.asarray(data['pv'])

    # 无储能基准（纯购电：光伏不够的部分全买）
    base_b = np.maximum(0.0, load - pv)
    base_cost = float((price * base_b).sum())

    # 表1 指定时段（段号 61,73,85,97,109,121 -> 0-index 60,72,84,96,108,120）
    sel = [('10:00-10:10', 60), ('12:00-12:10', 72), ('14:00-14:10', 84),
           ('16:00-16:10', 96), ('18:00-18:10', 108), ('20:00-20:10', 120)]

    # 表2 六个 4 小时块
    blocks_name = ['0:00-4:00', '4:00-8:00', '8:00-12:00',
                   '12:00-16:00', '16:00-20:00', '20:00-24:00']
    blocks = [(0, 24), (24, 48), (48, 72), (72, 96), (96, 120), (120, 144)]

    metrics = {
        '全天购电量_kWh': round(sol['total_purchase'], 4),
        '全天购电费_元': round(sol['cost'], 4),
        '无储能基准费用_元': round(base_cost, 4),
        '节省费用_元': round(base_cost - sol['cost'], 4),
        '节省比例_%': round(100 * (base_cost - sol['cost']) / base_cost, 3),
        '全天充电量_kWh': round(sol['total_charge'], 4),
        '全天放电量_kWh': round(sol['total_discharge'], 4),
        '峰值充电功率_kW': round(float(np.max(sol['x']) * 6), 2),
        '峰值放电功率_kW': round(float(np.max(sol['y']) * 6), 2),
        'SOC_min_kWh': round(sol['soc_min'], 4),
        'SOC_max_kWh': round(sol['soc_max'], 4),
        '表1_指定时段购电量_kWh': {name: round(float(sol['b'][i]), 4) for name, i in sel},
        '表2_充放电量_kWh': {
            name: {'充电量': round(float(sol['x'][a:z].sum()), 4),
                   '放电量': round(float(sol['y'][a:z].sum()), 4)}
            for name, (a, z) in zip(blocks_name, blocks)
        },
    }

    with open(os.path.join(CODE, 'metrics.json'), 'w', encoding='utf-8') as f:
        json.dump(metrics, f, ensure_ascii=False, indent=2)

    print('求解完成')
    print(f'全天购电量 = {metrics["全天购电量_kWh"]} kWh')
    print(f'全天购电费 = {metrics["全天购电费_元"]} 元')
    print(f'无储能基准 = {metrics["无储能基准费用_元"]} 元，节省 {metrics["节省费用_元"]} 元 ({metrics["节省比例_%"]}%)')
    print(f'SOC 范围 = [{metrics["SOC_min_kWh"]}, {metrics["SOC_max_kWh"]}] kWh')
    print('结果已写入', RESULT)


if __name__ == '__main__':
    main()
