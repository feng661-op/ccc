# -*- coding: utf-8 -*-
"""问题2结果独立验收：检查result2结构、数值对账、SOC/功率约束、防泄漏和紧急购电汇总。"""
from pathlib import Path
import csv
import json
import numpy as np
import openpyxl

HERE = Path(__file__).resolve().parent
Q2 = HERE.parent
XMAX = 5000.0 / 6.0
SOC_MIN, SOC_MAX = 1200.0, 10800.0


def main():
    z = np.load(HERE / 'run_data.npz')
    wb = openpyxl.load_workbook(Q2 / 'result2.xlsx', read_only=True, data_only=True)

    # 计划购电量
    rows = list(wb['计划购电量'].iter_rows(values_only=True))
    header, body = list(rows[0]), rows[1:]
    dates = [r[0].date().isoformat() for r in body]
    plan = np.asarray([[float(v or 0) for v in r[1:145]] for r in body], dtype=float)
    total = np.asarray([float(r[145] or 0) for r in body], dtype=float)
    fee = np.asarray([float(r[146] or 0) for r in body], dtype=float)

    # 充放电量
    crows = list(wb['充放电量'].iter_rows(values_only=True))
    cbody = crows[1:]
    cdates = [r[0].date().isoformat() for r in cbody if r[0] is not None]
    max_charge_block_error = 0.0
    max_discharge_block_error = 0.0
    for k, day in enumerate(range(31, 365)):
        for b in range(6):
            r = cbody[k * 6 + b]
            sl = slice(b * 24, (b + 1) * 24)
            max_charge_block_error = max(max_charge_block_error,
                abs(float(r[2] or 0) - float(z['charge'][day, sl].sum())))
            max_discharge_block_error = max(max_discharge_block_error,
                abs(float(r[3] or 0) - float(z['discharge'][day, sl].sum())))

    # 紧急购电量：模板展开后每个日期仅首事件行写日期，其余连续事件日期留空。
    erows = list(wb['紧急购电量'].iter_rows(values_only=True))
    ebody = erows[1:]
    edates = [r[0].date().isoformat() for r in ebody if r[0] is not None]
    workbook_emergency_sum = float(sum(float(r[2] or 0) for r in ebody))
    array_emergency_sum = float(z['emergency'][31:].sum())

    wb.close()

    # 防泄漏逐日审计
    with (HERE / 'leakage_audit.csv').open(encoding='utf-8-sig', newline='') as f:
        leakage = list(csv.DictReader(f))

    soc_min = float(min(z['soc00'].min(), z['soc24'].min()))
    soc_max = float(max(z['soc00'].max(), z['soc24'].max()))
    soc_cont = float(np.max(np.abs(z['soc24'][:-1] - z['soc00'][1:])))
    max_charge = float(z['charge'].max())
    max_discharge = float(z['discharge'].max())

    checks = {
        'plan_shape_334x144': len(body) == 334 and plan.shape == (334, 144),
        'plan_dates_exact': len(set(dates)) == 334 and dates[0] == '2025-02-01' and dates[-1] == '2025-12-31',
        'plan_time_boundary_exact': header[1] == '0:10-0:20' and header[144] == '0:00-0:10+1',
        'plan_nonnegative_finite': bool(np.all(np.isfinite(plan)) and np.all(plan >= -1e-8)),
        'daily_total_reconciled': float(np.max(np.abs(total - plan.sum(axis=1)))) <= 1e-4,
        'plan_array_reconciled': float(np.max(np.abs(plan - z['plan_q'][31:]))) <= 1e-5,
        'fee_array_reconciled': float(np.max(np.abs(fee - z['plan_cost'][31:]))) <= 1e-5,
        'charge_sheet_full_334x6': len(cbody) == 334 * 6 and len(cdates) == 334 and len(set(cdates)) == 334,
        'charge_block_reconciled': max_charge_block_error <= 1e-5,
        'discharge_block_reconciled': max_discharge_block_error <= 1e-5,
        'emergency_all_dates_present': len(edates) == 334 and len(set(edates)) == 334,
        'emergency_sum_reconciled': abs(workbook_emergency_sum - array_emergency_sum) <= 1e-4,
        'soc_bounds': soc_min >= SOC_MIN - 1e-6 and soc_max <= SOC_MAX + 1e-6,
        'soc_cross_day_continuity': soc_cont <= 1e-8,
        'charge_power_bound': max_charge <= XMAX + 1e-6,
        'discharge_power_bound': max_discharge <= XMAX + 1e-6,
        'leakage_audit_365_pass': len(leakage) == 365 and all(r['leakage_check'] == 'PASS' for r in leakage),
    }
    report = {
        'status': 'PASS' if all(checks.values()) else 'FAIL',
        'checks': checks,
        'diagnostics': {
            'plan_total_max_abs_error_kwh': float(np.max(np.abs(total - plan.sum(axis=1)))),
            'plan_npz_max_abs_error_kwh': float(np.max(np.abs(plan - z['plan_q'][31:]))),
            'fee_npz_max_abs_error_yuan': float(np.max(np.abs(fee - z['plan_cost'][31:]))),
            'charge_block_max_abs_error_kwh': max_charge_block_error,
            'discharge_block_max_abs_error_kwh': max_discharge_block_error,
            'emergency_workbook_sum_kwh': workbook_emergency_sum,
            'emergency_array_sum_kwh': array_emergency_sum,
            'soc_min_kwh': soc_min,
            'soc_max_kwh': soc_max,
            'soc_continuity_max_abs_error_kwh': soc_cont,
            'max_charge_10min_kwh': max_charge,
            'max_discharge_10min_kwh': max_discharge,
        }
    }
    (HERE / 'validation_report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if report['status'] != 'PASS':
        raise SystemExit(1)


if __name__ == '__main__':
    main()
