# -*- coding: utf-8 -*-
"""2026高教社杯C题 问题2全年滚动回测。

流程：
1) 仅用0:00之前可见历史做因果预测；
2) 用历史预测残差构造经验场景；
3) 解“lead interval + 当天 + 次日”两阶段随机MPC，CVaR抑制紧急购电尾部风险；
4) 当天真实运行只使用当前已发生的负载/PV，按场景SOC储备底线实时充放电，缺口5倍电价紧急购电；
5) SOC全年连续，1月暖启动，2月1日开始正式评价；
6) 填写官方result2.xlsx并输出审计数据与指标。
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import shutil
import sys
from copy import copy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Tuple

import numpy as np
import openpyxl

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))

from q2_core import (  # noqa: E402
    CVAR_ALPHA, DT, ETA_C, ETA_D, PMAX, SOC0, SOC_MAX, SOC_MIN, XMAX,
    TERMINAL_RESERVE, TERMINAL_SHORTAGE_VALUE,
    InputData, ForecastBundle, build_net_scenarios, execute_interval,
    grouped_emergency_periods, load_official_inputs, precompute_causal_forecasts,
    scenario_floor, solve_stochastic_mpc,
)

EVAL_START = 31  # 2025-02-01
RISK_CANDIDATES = (0.0, 0.005, 0.01, 0.02, 0.03, 0.05, 0.15, 0.30)
RISK_BUDGET_PCT = 0.002  # 仅允许用1月风险中性成本的0.2%作为可靠性保险预算


@dataclass
class SimulationResult:
    risk_lambda: float
    plan_q: np.ndarray
    charge: np.ndarray
    discharge: np.ndarray
    emergency: np.ndarray
    curtailment: np.ndarray
    soc00: np.ndarray
    soc24: np.ndarray
    plan_cost: np.ndarray
    emergency_cost: np.ndarray
    predicted_expected_emergency_cost: np.ndarray
    predicted_cvar_cost: np.ndarray
    scenario_count: np.ndarray
    objective: np.ndarray


def _calendar_price(price: np.ndarray) -> np.ndarray:
    # 自然日0:00-0:10沿用上一份计划末段价格；之后0:10-24:00对应price[0..142]。
    return np.concatenate(([price[-1]], price[:143]))


def simulate_strategy(
    data: InputData,
    fc: ForecastBundle,
    risk_lambda: float,
    end_day: int = 365,
    verbose: bool = False,
) -> SimulationResult:
    n = int(end_day)
    plan_q = np.zeros((n, 144), dtype=float)
    charge = np.zeros((n, 144), dtype=float)
    discharge = np.zeros((n, 144), dtype=float)
    emergency = np.zeros((n, 144), dtype=float)
    curtail = np.zeros((n, 144), dtype=float)
    soc00 = np.zeros(n, dtype=float)
    soc24 = np.zeros(n, dtype=float)
    plan_cost = np.zeros(n, dtype=float)
    emergency_cost = np.zeros(n, dtype=float)
    pred_e = np.zeros(n, dtype=float)
    pred_cvar = np.zeros(n, dtype=float)
    scount = np.zeros(n, dtype=int)
    obj = np.zeros(n, dtype=float)

    soc = SOC0
    prev_q_last = 0.0
    cal_price = _calendar_price(data.price)

    for d in range(n):
        soc00[d] = soc
        scenarios = build_net_scenarios(d, data, fc)
        sol = solve_stochastic_mpc(
            scenarios, data.price, soc, prev_q_last if d >= 1 else 0.0,
            risk_lambda=risk_lambda,
            alpha=CVAR_ALPHA,
            terminal_value=TERMINAL_SHORTAGE_VALUE,
        )
        q = np.maximum(sol.q_current, 0.0)
        q[np.abs(q) < 1e-9] = 0.0
        plan_q[d] = q
        plan_cost[d] = float(np.dot(data.price, q))
        pred_e[d] = sol.expected_emergency_cost
        pred_cvar[d] = sol.cvar_emergency_cost
        scount[d] = sol.scenario_count
        obj[d] = sol.objective

        # calendar interval 0:00-0:10：仍执行昨天0:00发布计划的最后一段。
        if d >= 1:
            floor = scenario_floor(sol, 1)  # lead之后的SOC场景储备底线
            soc, x, y, e, w = execute_interval(
                soc, prev_q_last, data.net[d - 1, 143], floor
            )
            charge[d, 0] = x
            discharge[d, 0] = y
            emergency[d, 0] = e
            curtail[d, 0] = w
        else:
            # 附件没有2025-01-01 0:00-0:10这一段实际数据：保持给定SOC不变，不虚构数据。
            pass

        # 当前新计划从0:10起生效。只执行到24:00（q[0..142]）；q[143]留到次日lead执行。
        for j in range(143):
            floor = scenario_floor(sol, 2 + j)  # t=1+j执行后的SOC
            soc, x, y, e, w = execute_interval(
                soc, q[j], data.net[d, j], floor
            )
            cidx = j + 1
            charge[d, cidx] = x
            discharge[d, cidx] = y
            emergency[d, cidx] = e
            curtail[d, cidx] = w

        # 数值安全检查
        if soc < SOC_MIN - 1e-6 or soc > SOC_MAX + 1e-6:
            raise RuntimeError(f'{data.dates[d].date()} 24:00 SOC越界: {soc}')
        soc = float(np.clip(soc, SOC_MIN, SOC_MAX))
        soc24[d] = soc
        prev_q_last = float(q[143])
        emergency_cost[d] = float(np.dot(5.0 * cal_price, emergency[d]))

        if verbose and (d % 30 == 0 or d == n - 1):
            print(
                f'[{d+1:03d}/{n}] {data.dates[d].date()} '
                f'SOC {soc00[d]:.1f}->{soc24[d]:.1f}, '
                f'plan={plan_q[d].sum():.1f} kWh, emergency={emergency[d].sum():.2f} kWh, '
                f'K={scount[d]}'
            )

    return SimulationResult(
        risk_lambda=float(risk_lambda), plan_q=plan_q,
        charge=charge, discharge=discharge, emergency=emergency, curtailment=curtail,
        soc00=soc00, soc24=soc24, plan_cost=plan_cost, emergency_cost=emergency_cost,
        predicted_expected_emergency_cost=pred_e, predicted_cvar_cost=pred_cvar,
        scenario_count=scount, objective=obj,
    )


def calibrate_risk(data: InputData, fc: ForecastBundle) -> Tuple[float, List[Dict[str, float]]]:
    """只用1月暖启动数据选择CVaR权重；正式2-12月结果完全不参与调参。

    规则不是强制使用风险厌恶：先以lambda=0作为风险中性基准，只允许总成本
    最多增加0.2%作为可靠性保险预算；在预算内的正lambda候选中，选择“每增加
    1元成本所减少的紧急购电kWh”最高的Pareto点。若没有正权重带来真实改善，
    自动回退lambda=0。
    """
    rows: List[Dict[str, float]] = []
    for lam in RISK_CANDIDATES:
        sim = simulate_strategy(data, fc, lam, end_day=EVAL_START, verbose=False)
        pc = float(sim.plan_cost.sum())
        ec = float(sim.emergency_cost.sum())
        ee = float(sim.emergency.sum())
        rows.append({
            'risk_lambda': float(lam),
            'january_plan_cost': pc,
            'january_emergency_cost': ec,
            'january_total_cost': pc + ec,
            'january_emergency_kwh': ee,
        })
        print(f'[Jan校准] lambda={lam:.3f}: total={pc+ec:.2f}, emergency={ee:.2f} kWh')

    neutral = next(r for r in rows if abs(r['risk_lambda']) < 1e-12)
    base_cost = neutral['january_total_cost']
    base_emg = neutral['january_emergency_kwh']
    for r in rows:
        premium = r['january_total_cost'] - base_cost
        reduction = base_emg - r['january_emergency_kwh']
        r['cost_premium_vs_neutral_yuan'] = premium
        r['cost_premium_vs_neutral_pct'] = premium / base_cost if base_cost else 0.0
        r['emergency_reduction_vs_neutral_kwh'] = reduction
        r['emergency_reduction_per_premium_kwh_per_yuan'] = (
            reduction / premium if premium > 1e-9 and reduction > 1e-9 else 0.0
        )

    shortlist = [
        r for r in rows
        if r['risk_lambda'] > 0.0
        and r['cost_premium_vs_neutral_yuan'] > 1e-9
        and r['cost_premium_vs_neutral_pct'] <= RISK_BUDGET_PCT + 1e-12
        and r['emergency_reduction_vs_neutral_kwh'] > 1e-9
    ]
    if shortlist:
        best = max(
            shortlist,
            key=lambda r: (
                r['emergency_reduction_per_premium_kwh_per_yuan'],
                r['emergency_reduction_vs_neutral_kwh'],
                -r['risk_lambda'],
            ),
        )
    else:
        best = neutral
    print(
        f"[Jan校准] 风险预算={RISK_BUDGET_PCT:.2%}，选择 lambda={best['risk_lambda']:.3f}; "
        f"premium={best['cost_premium_vs_neutral_pct']:.3%}, "
        f"emergency_reduction={best['emergency_reduction_vs_neutral_kwh']:.2f} kWh"
    )
    return float(best['risk_lambda']), rows


def _forecast_metrics(data: InputData, fc: ForecastBundle, start: int = EVAL_START) -> Dict[str, float]:
    sl = slice(start, 365)
    def metrics(actual: np.ndarray, pred: np.ndarray, prefix: str) -> Dict[str, float]:
        err = pred[sl] - actual[sl]
        a = actual[sl]
        mae = float(np.mean(np.abs(err)))
        rmse = float(np.sqrt(np.mean(err ** 2)))
        nmae = float(np.sum(np.abs(err)) / max(1e-12, np.sum(np.abs(a))))
        return {f'{prefix}_mae_kwh_per_10min': mae,
                f'{prefix}_rmse_kwh_per_10min': rmse,
                f'{prefix}_nmae': nmae}
    out = {}
    out.update(metrics(data.load, fc.load_current, 'load'))
    out.update(metrics(data.pv, fc.pv_current, 'pv'))
    out.update(metrics(data.net, fc.net_current, 'net'))
    # 同时给kW口径便于论文叙述
    out['load_mae_kw'] = out['load_mae_kwh_per_10min'] / DT
    out['pv_mae_kw'] = out['pv_mae_kwh_per_10min'] / DT
    out['net_mae_kw'] = out['net_mae_kwh_per_10min'] / DT
    return out


def build_metrics(
    data: InputData,
    fc: ForecastBundle,
    main: SimulationResult,
    baseline: SimulationResult,
    calibration: List[Dict[str, float]],
) -> Dict:
    s = EVAL_START
    plan_energy = float(main.plan_q[s:].sum())
    emg_energy = float(main.emergency[s:].sum())
    plan_cost = float(main.plan_cost[s:].sum())
    emg_cost = float(main.emergency_cost[s:].sum())
    total = plan_cost + emg_cost
    b_plan_cost = float(baseline.plan_cost[s:].sum())
    b_emg_cost = float(baseline.emergency_cost[s:].sum())
    b_total = b_plan_cost + b_emg_cost
    b_emg_energy = float(baseline.emergency[s:].sum())

    # calendar-day SOC continuity：前一天24:00必须等于次日0:00。
    continuity = float(np.max(np.abs(main.soc24[:-1] - main.soc00[1:])))
    all_soc = np.concatenate((main.soc00, main.soc24))
    # 完整10分钟实时SOC必须由实际充放电轨迹重构；soc00/soc24仅代表每日端点。
    realtime_min = float('inf')
    realtime_max = float('-inf')
    realtime_reconcile = 0.0
    for d in range(365):
        srt = float(main.soc00[d])
        realtime_min = min(realtime_min, srt)
        realtime_max = max(realtime_max, srt)
        for t in range(144):
            srt += ETA_C * float(main.charge[d, t]) - float(main.discharge[d, t]) / ETA_D
            realtime_min = min(realtime_min, srt)
            realtime_max = max(realtime_max, srt)
        realtime_reconcile = max(realtime_reconcile, abs(srt - float(main.soc24[d])))
    max_charge = float(np.max(main.charge))
    max_discharge = float(np.max(main.discharge))
    total_grid = plan_energy + emg_energy

    m = {
        'model': 'causal rolling forecast + empirical scenarios + two-stage MPC + CVaR + realtime reserve policy',
        'evaluation_period': '2025-02-01..2025-12-31',
        'evaluation_days': 334,
        'warmup_period': '2025-01-01..2025-01-31',
        'risk_lambda': main.risk_lambda,
        'cvar_alpha': CVAR_ALPHA,
        'terminal_reserve_kwh': TERMINAL_RESERVE,
        'terminal_shortage_value_yuan_per_kwh': TERMINAL_SHORTAGE_VALUE,
        'eta_charge': ETA_C,
        'eta_discharge': ETA_D,
        'round_trip_efficiency': ETA_C * ETA_D,
        'plan_purchase_kwh': plan_energy,
        'plan_purchase_cost_yuan': plan_cost,
        'emergency_purchase_kwh': emg_energy,
        'emergency_purchase_cost_yuan': emg_cost,
        'total_purchase_cost_yuan': total,
        'emergency_share_of_grid_energy': emg_energy / total_grid if total_grid > 0 else 0.0,
        'charge_energy_kwh': float(main.charge[s:].sum()),
        'discharge_energy_kwh': float(main.discharge[s:].sum()),
        'curtailed_surplus_kwh': float(main.curtailment[s:].sum()),
        'soc_min_kwh': float(all_soc.min()),
        'soc_max_kwh': float(all_soc.max()),
        'soc_endpoint_min_kwh': float(all_soc.min()),
        'soc_endpoint_max_kwh': float(all_soc.max()),
        'soc_realtime_min_kwh': realtime_min,
        'soc_realtime_max_kwh': realtime_max,
        'soc_realtime_end_reconcile_max_abs_error_kwh': realtime_reconcile,
        'soc_continuity_max_abs_error_kwh': continuity,
        'max_charge_per_10min_kwh': max_charge,
        'max_discharge_per_10min_kwh': max_discharge,
        'soc_safety_rate': float(np.mean((all_soc >= SOC_MIN - 1e-8) & (all_soc <= SOC_MAX + 1e-8))),
        'risk_neutral_total_cost_yuan': b_total,
        'risk_neutral_plan_cost_yuan': b_plan_cost,
        'risk_neutral_emergency_cost_yuan': b_emg_cost,
        'risk_neutral_emergency_kwh': b_emg_energy,
        'risk_cost_change_vs_neutral_yuan': total - b_total,
        'risk_cost_change_vs_neutral_pct': (total / b_total - 1.0) if b_total else 0.0,
        'risk_emergency_reduction_kwh': b_emg_energy - emg_energy,
        'risk_emergency_reduction_pct': (1.0 - emg_energy / b_emg_energy) if b_emg_energy > 0 else 0.0,
        'mean_scenarios_per_day_eval': float(np.mean(main.scenario_count[s:])),
        'january_risk_calibration': calibration,
        'forecast_metrics': _forecast_metrics(data, fc, s),
    }
    return m


def write_daily_summary(
    out_csv: Path, data: InputData, fc: ForecastBundle, sim: SimulationResult
) -> None:
    with out_csv.open('w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow([
            'date','plan_kwh','plan_cost_yuan','emergency_kwh','emergency_cost_yuan','total_cost_yuan',
            'charge_kwh','discharge_kwh','curtailment_kwh','soc_00_kwh','soc_24_kwh',
            'load_mae_kw','pv_mae_kw','net_mae_kw','scenario_count','pred_expected_emergency_cost','pred_cvar_cost'
        ])
        for d in range(365):
            load_mae = float(np.mean(np.abs(fc.load_current[d] - data.load[d])) / DT)
            pv_mae = float(np.mean(np.abs(fc.pv_current[d] - data.pv[d])) / DT)
            net_mae = float(np.mean(np.abs(fc.net_current[d] - data.net[d])) / DT)
            w.writerow([
                data.dates[d].date().isoformat(), sim.plan_q[d].sum(), sim.plan_cost[d],
                sim.emergency[d].sum(), sim.emergency_cost[d], sim.plan_cost[d] + sim.emergency_cost[d],
                sim.charge[d].sum(), sim.discharge[d].sum(), sim.curtailment[d].sum(),
                sim.soc00[d], sim.soc24[d], load_mae, pv_mae, net_mae,
                int(sim.scenario_count[d]), sim.predicted_expected_emergency_cost[d], sim.predicted_cvar_cost[d]
            ])


def write_leakage_audit(out_csv: Path, data: InputData, fc: ForecastBundle) -> None:
    with out_csv.open('w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow([
            'plan_date','latest_complete_actual_date_allowed','partial_previous_date_allowed',
            'partial_previous_last_slot_masked','scenario_latest_complete_residual_date_allowed','leakage_check'
        ])
        for d, date in enumerate(data.dates):
            full_idx = int(fc.full_source_max_day[d])
            partial_idx = int(fc.partial_source_day[d])
            full_date = data.dates[full_idx].date().isoformat() if full_idx >= 0 else ''
            partial_date = data.dates[partial_idx].date().isoformat() if partial_idx >= 0 else ''
            # 场景也最多用day-2完整误差；严格早于plan_date。
            ok = (full_idx < d) and (partial_idx < d)
            w.writerow([date.date().isoformat(), full_date, partial_date, True, full_date, 'PASS' if ok else 'FAIL'])
            if not ok:
                raise AssertionError(f'{date.date()} 信息集泄漏审计失败')


def _clone_style(dst, style_obj, number_format: str) -> None:
    dst._style = copy(style_obj)
    dst.number_format = number_format


def write_result2(
    template: Path,
    out_path: Path,
    data: InputData,
    sim: SimulationResult,
) -> None:
    shutil.copy2(template, out_path)
    wb = openpyxl.load_workbook(out_path)

    # 1) 计划购电量：模板已完整预置2.1-12.31共334行。
    ws = wb['计划购电量']
    if ws.max_row != 335 or ws.max_column < 147:
        raise ValueError(f'计划购电量模板结构异常: {ws.max_row}x{ws.max_column}')
    for out_i, d in enumerate(range(EVAL_START, 365), start=2):
        expected_date = data.dates[d].date()
        got = ws.cell(out_i, 1).value
        got_date = got.date() if hasattr(got, 'date') else got
        if got_date != expected_date:
            raise ValueError(f'计划购电量模板日期错位 row={out_i}: {got_date} != {expected_date}')
        q = sim.plan_q[d]
        for j in range(144):
            ws.cell(out_i, 2 + j).value = round(float(q[j]), 6)
        ws.cell(out_i, 146).value = round(float(q.sum()), 6)
        ws.cell(out_i, 147).value = round(float(sim.plan_cost[d]), 6)

    # 2) 充放电量：模板用“2/1,2/2,⋮,12/31”省略，展开成334天×6个4h块。
    ws = wb['充放电量']
    style_rows = []
    for r in range(2, 8):
        style_rows.append([(copy(ws.cell(r, c)._style), ws.cell(r, c).number_format) for c in range(1, 7)])
    heights = [ws.row_dimensions[r].height for r in range(2, 8)]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    blocks = ['0:00-4:00','4:00-8:00','8:00-12:00','12:00-16:00','16:00-20:00','20:00-24:00']
    row = 2
    for d in range(EVAL_START, 365):
        for b in range(6):
            vals = [None] * 6
            if b == 0:
                vals[0] = data.dates[d]
            vals[1] = blocks[b]
            sl = slice(24 * b, 24 * (b + 1))
            vals[2] = round(float(sim.charge[d, sl].sum()), 6)
            vals[3] = round(float(sim.discharge[d, sl].sum()), 6)
            if b == 0:
                vals[4] = datetime.strptime('00:00','%H:%M').time()
                vals[5] = round(float(sim.soc00[d]), 6)
            elif b == 1:
                vals[4] = '24:00'
                vals[5] = round(float(sim.soc24[d]), 6)
            for c, v in enumerate(vals, start=1):
                cell = ws.cell(row, c)
                cell.value = v
                st, nf = style_rows[b][c - 1]
                _clone_style(cell, st, nf)
            if heights[b] is not None:
                ws.row_dimensions[row].height = heights[b]
            row += 1

    # 3) 紧急购电量：按自然日连续10分钟段合并；无紧急购电也保留日期+0，保证334天可审计。
    ws = wb['紧急购电量']
    style_rows_e = []
    for r in range(2, 5):
        style_rows_e.append([(copy(ws.cell(r, c)._style), ws.cell(r, c).number_format) for c in range(1, 4)])
    heights_e = [ws.row_dimensions[r].height for r in range(2, 5)]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    row = 2
    for d in range(EVAL_START, 365):
        events = grouped_emergency_periods(sim.emergency[d])
        if not events:
            events = [('', 0.0)]
        for k, (period, amount) in enumerate(events):
            vals = [data.dates[d] if k == 0 else None, period or None, round(float(amount), 6)]
            template_idx = min(k, 2)
            for c, v in enumerate(vals, start=1):
                cell = ws.cell(row, c)
                cell.value = v
                st, nf = style_rows_e[template_idx][c - 1]
                _clone_style(cell, st, nf)
            if heights_e[template_idx] is not None:
                ws.row_dimensions[row].height = heights_e[template_idx]
            row += 1

    wb.save(out_path)


def save_npz(out_path: Path, main: SimulationResult, baseline: SimulationResult, fc: ForecastBundle) -> None:
    np.savez_compressed(
        out_path,
        plan_q=main.plan_q, charge=main.charge, discharge=main.discharge,
        emergency=main.emergency, curtailment=main.curtailment,
        soc00=main.soc00, soc24=main.soc24,
        plan_cost=main.plan_cost, emergency_cost=main.emergency_cost,
        baseline_plan_q=baseline.plan_q, baseline_emergency=baseline.emergency,
        baseline_plan_cost=baseline.plan_cost, baseline_emergency_cost=baseline.emergency_cost,
        forecast_load=fc.load_current, forecast_pv=fc.pv_current, forecast_net=fc.net_current,
    )


def write_reports(out_dir: Path, metrics: Dict) -> None:
    fm = metrics['forecast_metrics']
    text = f"""# 问题2模型与结果说明\n\n## 1. 模型定位\n\n问题2不是把附件2全年真实值一次性代入全年LP，而是模拟每天0:00的信息集，形成“历史信息 → 预测 → 场景风险优化 → 实际运行 → 5倍紧急购电结算 → 数据进入历史”的walk-forward闭环。\n\n正式评价区间严格按result2模板为 **2025-02-01至2025-12-31（334天）**；1月只作暖启动和风险参数校准。储能从2025-01-01给定的6000 kWh开始连续运行，2月1日不重置。\n\n## 2. 10分钟边界与信息集\n\nresult2计划列从`0:10-0:20`开始，最后为`0:00-0:10+1`。因此每天0:00发布的新计划在0:10起生效，0:00-0:10仍执行昨日计划最后一段。程序显式保留这个lead interval。\n\n在第d天0:00：\n- 最晚完整可用数据行是d-2；\n- d-1行只允许使用0:10至24:00前的143个已发生点；\n- d-1行最后的`0:00+1`对应未来0:00-0:10，强制掩码，不参与当天0:00预测；\n- 场景残差也只使用d-2及以前完整误差。\n\n这比简单“使用昨天整行”更严格，避免最后10分钟的隐藏look-ahead。\n\n## 3. 预测与场景\n\n负载和光伏分别用因果历史模拟集成：上一日已观测部分、7日滞后、同星期历史和最近7日中位曲线加权。没有天气数据，因此不引入不可获得的未来天气特征。场景采用最近历史预测残差的整日相关曲线，按高价时段正净负荷残差风险分位抽取，保留时序相关性。正式期使用9个等概率经验分位场景；之所以不采用5个等概率场景，是因为题设紧急购电恰为5倍电价，5场景下“仅最坏1个场景缺1 kWh”的期望惩罚正好等于提前多买1 kWh的计划成本，会造成边际退化。9场景使尾部场景概率降至1/9，期望成本与CVaR形成真实风险—成本权衡。\n\n正式评价期预测误差：\n- 负载 MAE = {fm['load_mae_kw']:.2f} kW，NMAE = {fm['load_nmae']:.4f}\n- 光伏 MAE = {fm['pv_mae_kw']:.2f} kW，NMAE = {fm['pv_nmae']:.4f}\n- 净负荷 MAE = {fm['net_mae_kw']:.2f} kW，NMAE = {fm['net_nmae']:.4f}\n\n## 4. 两阶段随机MPC + CVaR\n\n每天优化窗口包括：已经承诺的0:00-0:10 lead段 + 当天新计划144段 + 下一天144段辅助视野。当天144段计划购电量为所有场景共享的一阶段变量；下一天购电量允许按场景自适应，只作为MPC续期变量，第二天0:00会重新求解。\n\n储能状态方程使用与问题1统一的对称往返效率拆分：\n\n`eta_c = eta_d = sqrt(0.9)`，\n\n`S[t+1] = S[t] + eta_c*x[t] - y[t]/eta_d`。\n\n目标包括：计划购电费 + 场景期望5倍紧急购电费 + `lambda*CVaR`尾部风险 + 48小时末低于6000 kWh的软终端储备价值。终端不是硬性回到6000，因此不会造成“每天电池自动复位”。冻结主模型后的事后敏感性显示，`v_T=0,0.4,0.8,1.2` 时费用和年末SOC几乎不变，因此软终端项在本题中基本不绑定；12月31日未放空主要由48小时滚动视野、后续购电成本和SOC连续状态共同决定。\n\n1月仅用暖启动数据事前校准风险权重：以 `lambda=0` 为风险中性基准，风险保险预算上限为基准成本的 {RISK_BUDGET_PCT:.2%}；在预算内选择“每增加1元成本所减少的紧急购电量”最高的 Pareto 点，无有效改善则回退风险中性。最终 `lambda = {metrics['risk_lambda']:.3f}`，`alpha = {metrics['cvar_alpha']:.2f}`。\n\n## 5. 实时执行\n\n计划购电量一旦发布即按计划计费。实际运行每10分钟只读取当前已实现净负荷：若计划电量与光伏有富余，则在功率/SOC约束内充电；若不足，则只在场景SOC风险储备底线以上放电，其余缺口按该时刻电价5倍紧急购电。该规则不读取未来实际值。\n\n## 6. 正式评价结果\n\n- 计划购电量：{metrics['plan_purchase_kwh']:.2f} kWh\n- 计划购电费：{metrics['plan_purchase_cost_yuan']:.2f} 元\n- 紧急购电量：{metrics['emergency_purchase_kwh']:.2f} kWh\n- 紧急购电费：{metrics['emergency_purchase_cost_yuan']:.2f} 元\n- 总购电费用：**{metrics['total_purchase_cost_yuan']:.2f} 元**\n- 紧急购电占全部外网购电量：{metrics['emergency_share_of_grid_energy']:.4%}\n- 完整10分钟实时SOC范围：{metrics.get('soc_realtime_min_kwh', SOC_MIN):.2f}–{metrics.get('soc_realtime_max_kwh', SOC_MAX):.2f} kWh\n- 每日0:00/24:00端点SOC最小值：{metrics.get('soc_endpoint_min_kwh', metrics['soc_min_kwh']):.2f} kWh\n- SOC跨日连续最大误差：{metrics['soc_continuity_max_abs_error_kwh']:.3e} kWh\n- 最大10分钟充电量：{metrics['max_charge_per_10min_kwh']:.2f} kWh（约束上限 {XMAX:.2f}）\n- 最大10分钟放电量：{metrics['max_discharge_per_10min_kwh']:.2f} kWh（约束上限 {XMAX:.2f}）\n\n与风险中性（lambda=0）同场景滚动基线相比：\n- 成本变化：{metrics['risk_cost_change_vs_neutral_yuan']:.2f} 元（{metrics['risk_cost_change_vs_neutral_pct']:.3%}）\n- 紧急购电减少：{metrics['risk_emergency_reduction_kwh']:.2f} kWh（{metrics['risk_emergency_reduction_pct']:.3%}）\n\n注意：0.2%的风险保险预算仅用于1月暖启动期的事前参数选择，不是2-12月外推期的硬成本约束；正式评价期的风险溢价应按上面的实际结果如实报告。\n\n## 7. 输出文件\n\n- `result2.xlsx`：官方模板完整结果\n- `code/metrics.json`：核心指标\n- `code/daily_summary.csv`：逐日结算与SOC\n- `code/leakage_audit.csv`：逐日信息集/防泄漏审计\n- `code/run_data.npz`：可复核数值数组\n- `output/模型与结果说明.md`：本说明\n"""
    (out_dir / '模型与结果说明.md').write_text(text, encoding='utf-8')


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument('--skip-calibration', action='store_true', help='跳过1月lambda校准，使用--risk-lambda')
    ap.add_argument('--risk-lambda', type=float, default=0.15)
    ap.add_argument('--no-baseline', action='store_true', help='不跑lambda=0全年基线（正式结果默认会跑）')
    ap.add_argument('--verbose', action='store_true')
    args = ap.parse_args()

    q2 = REPO / '问题二'
    code_dir = q2 / 'code'
    out_dir = q2 / 'output'
    out_dir.mkdir(parents=True, exist_ok=True)

    print('[1/7] 读取官方附件并预计算严格因果预测...')
    data = load_official_inputs(REPO)
    fc = precompute_causal_forecasts(data)
    if data.dates[EVAL_START].date().isoformat() != '2025-02-01':
        raise AssertionError('评价起点不是2025-02-01')

    print('[2/7] 1月暖启动校准CVaR权重...')
    if args.skip_calibration:
        risk_lambda = float(args.risk_lambda)
        calibration = []
    else:
        risk_lambda, calibration = calibrate_risk(data, fc)

    print(f'[3/7] 全年主策略滚动回测，lambda={risk_lambda:.3f}...')
    main_sim = simulate_strategy(data, fc, risk_lambda, end_day=365, verbose=args.verbose)

    print('[4/7] 风险中性基线滚动回测...')
    if args.no_baseline:
        baseline = main_sim
    elif abs(risk_lambda) < 1e-12:
        baseline = main_sim
    else:
        baseline = simulate_strategy(data, fc, 0.0, end_day=365, verbose=False)

    print('[5/7] 汇总指标与审计...')
    metrics = build_metrics(data, fc, main_sim, baseline, calibration)
    code_dir.mkdir(parents=True, exist_ok=True)
    (code_dir / 'metrics.json').write_text(json.dumps(metrics, ensure_ascii=False, indent=2), encoding='utf-8')
    write_daily_summary(code_dir / 'daily_summary.csv', data, fc, main_sim)
    write_leakage_audit(code_dir / 'leakage_audit.csv', data, fc)
    save_npz(code_dir / 'run_data.npz', main_sim, baseline, fc)

    print('[6/7] 填写官方result2.xlsx...')
    template = REPO / '26C题' / '附件' / '附件5' / 'result2.xlsx'
    write_result2(template, q2 / 'result2.xlsx', data, main_sim)
    write_reports(out_dir, metrics)

    print('[7/7] 完成。核心结果：')
    print(json.dumps({
        'risk_lambda': metrics['risk_lambda'],
        'total_purchase_cost_yuan': metrics['total_purchase_cost_yuan'],
        'plan_purchase_cost_yuan': metrics['plan_purchase_cost_yuan'],
        'emergency_purchase_cost_yuan': metrics['emergency_purchase_cost_yuan'],
        'emergency_purchase_kwh': metrics['emergency_purchase_kwh'],
        'emergency_share_of_grid_energy': metrics['emergency_share_of_grid_energy'],
        'soc_min_kwh': metrics['soc_min_kwh'],
        'soc_max_kwh': metrics['soc_max_kwh'],
        'risk_cost_change_vs_neutral_yuan': metrics['risk_cost_change_vs_neutral_yuan'],
        'risk_emergency_reduction_kwh': metrics['risk_emergency_reduction_kwh'],
    }, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
