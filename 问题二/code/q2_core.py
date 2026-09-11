# -*- coding: utf-8 -*-
"""2026 C题问题2：因果滚动预测 + 场景风险优化 + 实时储能修正核心模块。

关键时间约定
------------
附件2每行144个样本与result2“计划购电量”144列一一对应：
0:10-0:20, ..., 23:50-0:00+1, 0:00-0:10+1。
因此每天0:00发布的新计划从0:10开始生效；0:00-0:10仍执行前一天计划的最后一段。
本模块显式建模这个10分钟lead interval，避免把当天未来实际值泄漏给0:00决策。
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np
import openpyxl
from scipy.optimize import linprog
from scipy.sparse import lil_matrix

# 储能参数（与问题1统一）
CAP = 12000.0
PMAX = 5000.0
DT = 1.0 / 6.0
XMAX = PMAX * DT
SOC0 = 6000.0
SOC_MIN = 1200.0
SOC_MAX = 10800.0
ETA_C = float(np.sqrt(0.9))
ETA_D = float(np.sqrt(0.9))

# 风险/MPC参数
# 9个等概率经验分位场景：单个尾部场景概率1/9<20%。
# 这点很关键：紧急购电是5倍电价，若只用5个等概率场景，
# “最坏1个场景缺1 kWh”的期望紧急成本恰好=提前计划多买1 kWh的成本，
# 会造成风险中性目标与CVaR项的边际退化。9场景可形成真实的期望-尾部风险权衡。
SCENARIO_QUANTILES = (0.05, 0.15, 0.25, 0.35, 0.50, 0.65, 0.75, 0.85, 0.95)
SCENARIO_WINDOW_DAYS = 42
CVAR_ALPHA = 0.80
TERMINAL_RESERVE = 6000.0
TERMINAL_SHORTAGE_VALUE = 0.80  # 元/kWh，软终端储备缺口价值
REALTIME_SOC_FLOOR_QUANTILE = 0.25


@dataclass
class InputData:
    dates: List[datetime]
    price: np.ndarray                  # (144,), 元/kWh
    load_kw: np.ndarray                # (365,144)
    pv_kw: np.ndarray                  # (365,144)
    load: np.ndarray                   # kWh/10min
    pv: np.ndarray                     # kWh/10min
    net: np.ndarray                    # load-pv, kWh/10min
    sample_headers: List[str]
    plan_headers: List[str]


@dataclass
class ForecastBundle:
    load_current: np.ndarray           # (365,144)
    pv_current: np.ndarray
    net_current: np.ndarray
    load_future: np.ndarray            # 第d天00:00对d+1的预测
    pv_future: np.ndarray
    net_future: np.ndarray
    full_source_max_day: np.ndarray    # 每日预测允许使用的最后完整数据行，=d-2
    partial_source_day: np.ndarray     # =d-1，仅使用0..142


@dataclass
class DayPlanSolution:
    q_current: np.ndarray              # 144段当天新发布计划
    scenario_soc: np.ndarray           # (K,H+1), H=289（lead+当前144+未来144）
    scenario_emergency: np.ndarray     # (K,H)
    objective: float
    expected_emergency_cost: float
    cvar_emergency_cost: float
    scenario_count: int
    solve_status: str


def _as_time_label(v) -> str:
    if isinstance(v, str):
        return v
    if hasattr(v, 'strftime'):
        return v.strftime('%H:%M')
    return str(v)


def load_official_inputs(repo_root: Path) -> InputData:
    """读取附件1电价、附件2全年实际负载/PV，并检查144点结构。"""
    att = repo_root / '26C题' / '附件'
    p1 = att / '附件1.xlsx'
    p2 = att / '附件2.xlsx'
    tpl = att / '附件5' / 'result2.xlsx'

    wb1 = openpyxl.load_workbook(p1, read_only=True, data_only=True)
    ws1 = wb1['Sheet1']
    rows1 = list(ws1.iter_rows(min_row=2, values_only=True))
    wb1.close()
    if len(rows1) != 144:
        raise ValueError(f'附件1应有144个10分钟点，实际{len(rows1)}')
    price = np.asarray([float(r[1]) for r in rows1], dtype=float)
    sample_headers = [_as_time_label(r[0]) for r in rows1]

    wb2 = openpyxl.load_workbook(p2, read_only=True, data_only=True)
    wsl = wb2['小区负载']
    wsp = wb2['光伏发电实际功率']
    if wsl.max_row != 366 or wsp.max_row != 366 or wsl.max_column != 145 or wsp.max_column != 145:
        raise ValueError('附件2结构异常：预期365天×144点')
    dates: List[datetime] = []
    load_rows = list(wsl.iter_rows(min_row=2, max_row=366, values_only=True))
    pv_rows = list(wsp.iter_rows(min_row=2, max_row=366, values_only=True))
    if len(load_rows) != 365 or len(pv_rows) != 365:
        raise ValueError('附件2日期行数异常')
    load_kw = np.empty((365, 144), dtype=float)
    pv_kw = np.empty((365, 144), dtype=float)
    for i, (lr, pr) in enumerate(zip(load_rows, pv_rows)):
        dv = lr[0]
        if isinstance(dv, datetime):
            dates.append(dv)
        else:
            dates.append(datetime.fromisoformat(str(dv)))
        load_kw[i] = np.asarray([float(v or 0.0) for v in lr[1:145]], dtype=float)
        pv_kw[i] = np.asarray([float(v or 0.0) for v in pr[1:145]], dtype=float)
    wb2.close()

    wbt = openpyxl.load_workbook(tpl, read_only=True, data_only=True)
    wspq = wbt['计划购电量']
    plan_headers = [str(wspq.cell(1, c).value) for c in range(2, 146)]
    wbt.close()
    if len(plan_headers) != 144:
        raise ValueError('result2计划购电量模板时间列不是144列')

    load = load_kw * DT
    pv = pv_kw * DT
    net = load - pv
    return InputData(
        dates=dates, price=price, load_kw=load_kw, pv_kw=pv_kw,
        load=load, pv=pv, net=net,
        sample_headers=sample_headers, plan_headers=plan_headers,
    )


def _target_date(start: datetime, target_day: int) -> datetime:
    return start + timedelta(days=int(target_day))


def _masked_previous_profile(actual: np.ndarray, plan_day: int) -> Optional[np.ndarray]:
    """在plan_day 0:00可见的上一行：0..142已发生；143(0:00-0:10)仍未来，必须屏蔽。"""
    if plan_day <= 0:
        return None
    p = actual[plan_day - 1].copy()
    # last slot must never use the future actual value
    if plan_day >= 2:
        lo = max(0, plan_day - 15)
        hi = plan_day - 1  # 不含上一日
        hist = actual[lo:hi, 143]
        fill = float(np.median(hist)) if hist.size else float(p[142])
    else:
        fill = float(p[142])
    p[143] = fill
    return p


def causal_profile_forecast(
    actual: np.ndarray,
    dates: Sequence[datetime],
    plan_day: int,
    target_day: int,
    kind: str,
) -> np.ndarray:
    """只使用plan_day 0:00之前可获得的信息预测target_day完整144点。

    完整历史最多到 plan_day-2；plan_day-1 只允许0..142，最后10分钟强制掩码。
    target_day可为plan_day或plan_day+1（扩展MPC视野）。
    """
    n_days = actual.shape[0]
    full_end = plan_day - 2
    start = dates[0]
    td = _target_date(start, target_day)
    prev_masked = _masked_previous_profile(actual, plan_day)

    candidates: List[np.ndarray] = []
    weights: List[float] = []
    is_current = (target_day == plan_day)

    if prev_masked is not None:
        candidates.append(prev_masked)
        weights.append(0.35 if (kind == 'load' and is_current) else
                       0.30 if (kind == 'pv' and is_current) else
                       0.20)

    lag7 = target_day - 7
    if 0 <= lag7 <= full_end and lag7 < n_days:
        candidates.append(actual[lag7])
        weights.append(0.35 if kind == 'load' else 0.25)

    if full_end >= 0:
        full_idx = list(range(0, min(full_end, n_days - 1) + 1))
        same_weekday = [j for j in full_idx if dates[j].weekday() == td.weekday()][-4:]
        if same_weekday:
            candidates.append(np.mean(actual[same_weekday], axis=0))
            weights.append(0.20 if kind == 'load' else 0.20)

        recent = full_idx[-7:]
        if recent:
            candidates.append(np.median(actual[recent], axis=0))
            weights.append(0.10 if kind == 'load' else 0.25)

    if not candidates:
        if kind == 'load':
            # 1月1日无历史：保守中性初始先验，仅影响暖启动；随后立即被真实历史替代。
            return np.full(144, 5000.0 * DT, dtype=float)
        return np.zeros(144, dtype=float)

    w = np.asarray(weights, dtype=float)
    w /= w.sum()
    pred = np.zeros(144, dtype=float)
    for wi, ci in zip(w, candidates):
        pred += wi * ci
    pred = np.maximum(pred, 0.0)
    return pred


def precompute_causal_forecasts(data: InputData) -> ForecastBundle:
    n = len(data.dates)
    lc = np.zeros((n, 144), dtype=float)
    pc = np.zeros((n, 144), dtype=float)
    lf = np.zeros((n, 144), dtype=float)
    pf = np.zeros((n, 144), dtype=float)
    max_full = np.full(n, -1, dtype=int)
    partial = np.full(n, -1, dtype=int)
    for d in range(n):
        lc[d] = causal_profile_forecast(data.load, data.dates, d, d, 'load')
        pc[d] = causal_profile_forecast(data.pv, data.dates, d, d, 'pv')
        lf[d] = causal_profile_forecast(data.load, data.dates, d, d + 1, 'load')
        pf[d] = causal_profile_forecast(data.pv, data.dates, d, d + 1, 'pv')
        max_full[d] = d - 2
        partial[d] = d - 1 if d >= 1 else -1
    return ForecastBundle(
        load_current=lc, pv_current=pc, net_current=lc - pc,
        load_future=lf, pv_future=pf, net_future=lf - pf,
        full_source_max_day=max_full, partial_source_day=partial,
    )


def _historical_residuals(
    day: int,
    actual_net: np.ndarray,
    forecast_net: np.ndarray,
) -> Dict[int, np.ndarray]:
    # 在day 0:00，完整可用于误差学习的数据行最多为day-2。
    end = day - 2
    if end < 0:
        return {}
    return {j: actual_net[j] - forecast_net[j] for j in range(end + 1)}


def build_net_scenarios(
    day: int,
    data: InputData,
    forecasts: ForecastBundle,
) -> np.ndarray:
    """构造lead+当前日+下一日的经验残差场景，形状(K,289)。

    只从 day-2 及更早的完整预测误差中取样，严格避免当天/上一行最后10分钟泄漏。
    """
    lead_point = 0.0 if day == 0 else float(forecasts.net_current[day - 1, 143])
    point = np.concatenate((
        np.asarray([lead_point], dtype=float),
        forecasts.net_current[day],
        forecasts.net_future[day],
    ))
    residuals = _historical_residuals(day, data.net, forecasts.net_current)

    # 需要历史三联：j-1的末10分钟作为lead误差，j为当前日误差，j+1为未来日误差。
    centers = [j for j in range(1, day - 2) if (j - 1 in residuals and j in residuals and j + 1 in residuals)]
    if not centers:
        return point[None, :]
    centers = centers[-SCENARIO_WINDOW_DAYS:]

    # 按“高价时段正残差风险”给历史情景排序，再取经验分位，保留相关的整日误差形状。
    scores = []
    price = data.price
    for j in centers:
        r = residuals[j]
        score = float(np.dot(price, np.maximum(r, 0.0)))
        scores.append(score)
    order = np.argsort(scores)
    ranked = [centers[int(i)] for i in order]

    picked: List[int] = []
    m = len(ranked)
    for q in SCENARIO_QUANTILES:
        idx = int(round(q * (m - 1)))
        j = ranked[idx]
        if j not in picked:
            picked.append(j)
    # 样本过少时把未选的最近情景补齐；不重复虚构。
    for j in reversed(centers):
        if len(picked) >= min(len(SCENARIO_QUANTILES), len(centers)):
            break
        if j not in picked:
            picked.append(j)

    scenarios = []
    for j in picked:
        residual_vec = np.concatenate((
            np.asarray([residuals[j - 1][143]], dtype=float),
            residuals[j],
            residuals[j + 1],
        ))
        s = point + residual_vec
        # 留宽裕物理边界，避免极端残差平移制造完全不合理场景；不截断正常负净负荷。
        s = np.clip(s, -2000.0, 2500.0)
        scenarios.append(s)
    return np.asarray(scenarios, dtype=float)


def solve_stochastic_mpc(
    scenario_net: np.ndarray,
    price: np.ndarray,
    soc0: float,
    q_lead: float,
    risk_lambda: float,
    alpha: float = CVAR_ALPHA,
    terminal_value: float = TERMINAL_SHORTAGE_VALUE,
) -> DayPlanSolution:
    """两阶段随机LP：当天144段q为共同一阶段决策，下一日q为场景自适应的MPC辅助决策。

    H=289 = 1个已承诺lead interval + 当天144段 + 下一天144段。
    lead段购电量已由昨天承诺，今天不可修改且不重复计计划购电费。
    """
    scenario_net = np.asarray(scenario_net, dtype=float)
    if scenario_net.ndim != 2 or scenario_net.shape[1] != 289:
        raise ValueError(f'scenario_net应为(K,289)，实际{scenario_net.shape}')
    K, H = scenario_net.shape
    if K < 1:
        raise ValueError('至少需要1个场景')
    price = np.asarray(price, dtype=float)
    if price.shape != (144,):
        raise ValueError('price必须144点')

    nq = 144
    nqf = 144
    block = nqf + 3 * H + (H + 1) + 1  # qfuture,x,y,e,s,r_terminal
    q0 = 0
    scen0 = nq
    zeta_idx = scen0 + K * block
    u0 = zeta_idx + 1
    N = u0 + K

    def offs(k: int):
        b = scen0 + k * block
        qf = b
        x = qf + nqf
        y = x + H
        e = y + H
        s = e + H
        r = s + H + 1
        return qf, x, y, e, s, r

    c = np.zeros(N, dtype=float)
    c[q0:q0 + nq] = price
    p_h = np.concatenate(([price[-1]], price, price))
    prob = 1.0 / K
    for k in range(K):
        qf, x, y, e, s, r = offs(k)
        c[qf:qf + nqf] = prob * price
        c[e:e + H] = prob * 5.0 * p_h
        c[r] = prob * terminal_value
    if risk_lambda > 0:
        c[zeta_idx] = risk_lambda
        c[u0:u0 + K] = risk_lambda / ((1.0 - alpha) * K)

    # 等式：每场景固定起始SOC + H段SOC递推
    n_eq = K * (H + 1)
    Aeq = lil_matrix((n_eq, N), dtype=float)
    beq = np.zeros(n_eq, dtype=float)
    rr = 0
    for k in range(K):
        qf, x, y, e, s, rterm = offs(k)
        Aeq[rr, s] = 1.0
        beq[rr] = soc0
        rr += 1
        for t in range(H):
            Aeq[rr, s + t + 1] = 1.0
            Aeq[rr, s + t] = -1.0
            Aeq[rr, x + t] = -ETA_C
            Aeq[rr, y + t] = 1.0 / ETA_D
            rr += 1

    # 不等式：平衡 + 软终端储备 + CVaR
    n_ub = K * H + K + K
    Aub = lil_matrix((n_ub, N), dtype=float)
    bub = np.zeros(n_ub, dtype=float)
    rr = 0
    for k in range(K):
        qf, x, y, e, s, rterm = offs(k)
        for t in range(H):
            # q + y + emergency >= net + x => x-y-e-q <= -net
            Aub[rr, x + t] = 1.0
            Aub[rr, y + t] = -1.0
            Aub[rr, e + t] = -1.0
            if t == 0:
                bub[rr] = float(q_lead - scenario_net[k, t])
            elif 1 <= t <= 144:
                Aub[rr, q0 + (t - 1)] = -1.0
                bub[rr] = -float(scenario_net[k, t])
            else:
                Aub[rr, qf + (t - 145)] = -1.0
                bub[rr] = -float(scenario_net[k, t])
            rr += 1
        # r_terminal >= TERMINAL_RESERVE - S_H
        Aub[rr, s + H] = -1.0
        Aub[rr, rterm] = -1.0
        bub[rr] = -TERMINAL_RESERVE
        rr += 1

    # CVaR：Z_k - zeta - u_k <= 0，Z只针对lead+当前承诺日(t=0..144)紧急购电风险。
    for k in range(K):
        qf, x, y, e, s, rterm = offs(k)
        for t in range(145):
            Aub[rr, e + t] = 5.0 * p_h[t]
        Aub[rr, zeta_idx] = -1.0
        Aub[rr, u0 + k] = -1.0
        bub[rr] = 0.0
        rr += 1

    bounds: List[Tuple[float, Optional[float]]] = [(0.0, None)] * N
    # 精确覆盖各变量边界
    for k in range(K):
        qf, x, y, e, s, rterm = offs(k)
        for i in range(x, x + H):
            bounds[i] = (0.0, XMAX)
        for i in range(y, y + H):
            bounds[i] = (0.0, XMAX)
        for i in range(e, e + H):
            bounds[i] = (0.0, None)
        for i in range(s, s + H + 1):
            bounds[i] = (SOC_MIN, SOC_MAX)
        bounds[rterm] = (0.0, None)
    bounds[zeta_idx] = (0.0, None)
    for i in range(u0, u0 + K):
        bounds[i] = (0.0, None)

    res = linprog(
        c,
        A_ub=Aub.tocsr(), b_ub=bub,
        A_eq=Aeq.tocsr(), b_eq=beq,
        bounds=bounds,
        method='highs',
        options={'presolve': True},
    )
    if not res.success:
        raise RuntimeError(f'随机MPC求解失败: {res.status} {res.message}')
    v = res.x
    q_current = v[q0:q0 + nq].copy()
    ss = np.zeros((K, H + 1), dtype=float)
    ee = np.zeros((K, H), dtype=float)
    scenario_costs = np.zeros(K, dtype=float)
    for k in range(K):
        qf, x, y, e, s, rterm = offs(k)
        ss[k] = v[s:s + H + 1]
        ee[k] = v[e:e + H]
        scenario_costs[k] = float(np.dot(5.0 * p_h[:145], ee[k, :145]))
    expected_emergency = float(np.mean(scenario_costs))
    sorted_costs = np.sort(scenario_costs)
    tail_n = max(1, int(np.ceil((1.0 - alpha) * K)))
    cvar = float(np.mean(sorted_costs[-tail_n:]))
    return DayPlanSolution(
        q_current=q_current,
        scenario_soc=ss,
        scenario_emergency=ee,
        objective=float(res.fun),
        expected_emergency_cost=expected_emergency,
        cvar_emergency_cost=cvar,
        scenario_count=K,
        solve_status=str(res.message),
    )


def execute_interval(
    soc: float,
    q: float,
    net_actual: float,
    soc_floor_after: float,
) -> Tuple[float, float, float, float, float]:
    """实时因果执行一段：先利用已承诺购电/PV余量充电；缺口时在风险储备底线以上放电，再紧急购电。

    返回(new_soc, charge, discharge, emergency, curtailed_surplus)。
    """
    margin = float(q - net_actual)  # q + PV - load
    if margin >= 0.0:
        headroom_grid = max(0.0, (SOC_MAX - soc) / ETA_C)
        charge = min(margin, XMAX, headroom_grid)
        new_soc = soc + ETA_C * charge
        curtail = max(0.0, margin - charge)
        return new_soc, charge, 0.0, 0.0, curtail

    deficit = -margin
    floor = float(np.clip(soc_floor_after, SOC_MIN, SOC_MAX))
    # 实际SOC若已低于场景底线，不再为了当前缺口继续透支；硬下限仍由SOC_MIN兜底。
    deliverable_by_floor = max(0.0, (soc - max(SOC_MIN, floor)) * ETA_D)
    discharge = min(deficit, XMAX, deliverable_by_floor)
    new_soc = soc - discharge / ETA_D
    emergency = max(0.0, deficit - discharge)
    return new_soc, 0.0, discharge, emergency, 0.0


def scenario_floor(solution: DayPlanSolution, state_index: int) -> float:
    vals = solution.scenario_soc[:, state_index]
    return float(np.quantile(vals, REALTIME_SOC_FLOOR_QUANTILE))


def natural_interval_label(i: int) -> str:
    """自然日第i个10分钟段，i=0 -> 0:00-0:10，i=143 -> 23:50-0:00+1。"""
    a = i * 10
    b = (i + 1) * 10
    def fmt(m: int, end: bool = False) -> str:
        if m == 1440:
            return '0:00+1' if end else '24:00'
        return f'{m // 60}:{m % 60:02d}'
    return f'{fmt(a)}-{fmt(b, end=True)}'


def grouped_emergency_periods(emergency: np.ndarray, tol: float = 1e-7) -> List[Tuple[str, float]]:
    """把连续发生紧急购电的10分钟段合并成时间段，购电量求和。"""
    e = np.asarray(emergency, dtype=float)
    idx = np.where(e > tol)[0]
    if idx.size == 0:
        return []
    groups: List[Tuple[int, int]] = []
    start = prev = int(idx[0])
    for x in idx[1:]:
        x = int(x)
        if x == prev + 1:
            prev = x
        else:
            groups.append((start, prev))
            start = prev = x
    groups.append((start, prev))

    out = []
    for a, z in groups:
        start_m = a * 10
        end_m = (z + 1) * 10
        def fmt(m: int, end: bool = False) -> str:
            if m == 1440:
                return '0:00+1' if end else '24:00'
            return f'{m // 60}:{m % 60:02d}'
        out.append((f'{fmt(start_m)}-{fmt(end_m, end=True)}', float(e[a:z + 1].sum())))
    return out
